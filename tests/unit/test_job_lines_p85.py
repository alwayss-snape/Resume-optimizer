"""P8.5: job-line formats outside tech: title / company / location / dates
read from one line or several, and nothing inherited from the job before."""
import docx
import pytest

from app.analysis.resume_normalizer import ResumeNormalizer
from app.ingestion.docx import DocxParser


def _jobs(lines, tmp_path, heading="Experience"):
    d = docx.Document()
    d.add_paragraph().add_run("Avery Lee").bold = True
    d.add_paragraph().add_run(heading.upper()).bold = True
    for kind, text in lines:
        if kind == "bold":
            d.add_paragraph().add_run(text).bold = True
        elif kind == "b":
            d.add_paragraph(text, style="List Bullet")
        else:
            d.add_paragraph(text)
    path = str(tmp_path / "r.docx")
    d.save(path)
    resume = ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume
    return [(e.title, e.company, e.location, e.start_date, e.end_date, len(e.bullets)) for e in resume.experience], resume


@pytest.mark.parametrize("line, want", [
    ("Shift Supervisor, Starbucks, Atlanta GA, 2021 - Present",
     ("Shift Supervisor", "Starbucks", "Atlanta GA", "2021", "Present")),
    ("Senior Accountant, Brightline Media, New York, NY\t07/2022 – Present",
     ("Senior Accountant", "Brightline Media", "New York, NY", "07/2022", "Present")),
    ("Logistics NCOIC (92Y Unit Supply Specialist), U.S. Army, Fort Hood, TX\t2014 – 2024",
     ("Logistics NCOIC (92Y Unit Supply Specialist)", "U.S. Army", "Fort Hood, TX", "2014", "2024")),
    ("Supply Chain Analyst, Groupe SEB, Lyon\t01/09/2019 – 31/08/2023",
     ("Supply Chain Analyst", "Groupe SEB", "Lyon", "01/09/2019", "31/08/2023")),  # a known city (P9.21)
    ("Senior Software Engineer, Infosys Ltd, Bengaluru\tJuly 2018 – Till Date",  # the live india run (P9.21)
     ("Senior Software Engineer", "Infosys Ltd", "Bengaluru", "July 2018", "Till Date")),
    ("Software Engineer, TCS, Chennai\tJune 2015 – June 2018",
     ("Software Engineer", "TCS", "Chennai", "June 2015", "June 2018")),
    ("Data Engineer, Acme Pvt Ltd, Coimbatore\t2019 – 2021",  # after a company suffix, a place-like name
     ("Data Engineer", "Acme Pvt Ltd", "Coimbatore", "2019", "2021")),
    ("Backend Engineer, Payments Platform, Stripe\t2019 – 2021",  # an unknown name stays with the company
     ("Backend Engineer", "Payments Platform, Stripe", None, "2019", "2021")),
    ("SVP, Global Supply Chain | Meridian Consumer Products | 2013 – 2018",
     ("SVP, Global Supply Chain", "Meridian Consumer Products", None, "2013", "2018")),
    ("Delivery Driver (part-time, concurrent) — DoorDash — 2020 to present",
     ("Delivery Driver (part-time, concurrent)", "DoorDash", None, "2020", "present")),
    ("Warehouse Associate — FedEx Ground Hub, Memphis TN — 2019 to present",
     ("Warehouse Associate", "FedEx Ground Hub", "Memphis TN", "2019", "present")),
    ("Judicial Law Clerk, Hon. Ellen Park, U.S. District Court for the District of Maryland\tAug 2017 – Aug 2018",
     ("Judicial Law Clerk", "Hon. Ellen Park, U.S. District Court for the District of Maryland", None,
      "Aug 2017", "Aug 2018")),
])
def test_one_line_job_formats(line, want, tmp_path):
    jobs, _ = _jobs([("bold", line), ("b", "Did the work.")], tmp_path)
    assert jobs == [(*want, 1)]


def test_title_tab_company_then_a_dates_line(tmp_path):
    jobs, _ = _jobs([("bold", "Registered Nurse - Step-Down Unit\tBanner University Medical Center, Phoenix, AZ"),
                     ("p", "03/2021 - Present"), ("b", "Provide direct patient care."),
                     ("bold", "Registered Nurse - Medical-Surgical\tSt. Joseph's Hospital, Phoenix, AZ"),
                     ("p", "06/2019 - 02/2021"), ("b", "Cared for post-operative patients.")], tmp_path)
    assert jobs == [
        ("Registered Nurse - Step-Down Unit", "Banner University Medical Center", "Phoenix, AZ", "03/2021", "Present", 1),
        ("Registered Nurse - Medical-Surgical", "St. Joseph's Hospital", "Phoenix, AZ", "06/2019", "02/2021", 1)]


def test_federal_header_over_three_lines(tmp_path):
    jobs, resume = _jobs([
        ("bold", "Management and Program Analyst (GS-0343-12)"), ("p", "U.S. Department of Agriculture, Washington, DC"),
        ("p", "10/2019 – Present | 40 hours per week | Salary: $94,199 per year | Supervisor: John Doe, may contact"),
        ("b", "Conduct program evaluations."), ("b", "Manage $3M budget execution.")], tmp_path, "Work Experience")
    assert jobs == [("Management and Program Analyst (GS-0343-12)", "U.S. Department of Agriculture",
                     "Washington, DC", "10/2019", "Present", 2)]
    assert resume.experience[0].details == [
        "40 hours per week | Salary: $94,199 per year | Supervisor: John Doe, may contact"]


def test_long_title_line_is_a_job_not_a_bullet(tmp_path):
    line = ("Judicial Law Clerk to the Honorable Ellen Park, United States District Court for the District of "
            "Maryland, Baltimore, MD\tAug 2017 – Aug 2018")
    assert len(line) > 100
    jobs, _ = _jobs([("bold", "Associate, Hale & Morgan LLP, Washington, DC\tSep 2018 – Present"),
                     ("b", "Draft motions."), ("bold", line), ("b", "Drafted bench memoranda.")], tmp_path)
    assert [j[0] for j in jobs] == ["Associate", "Judicial Law Clerk to the Honorable Ellen Park"]
    assert [j[5] for j in jobs] == [1, 1]


def test_later_role_never_inherits_a_date_column(tmp_path):
    jobs, _ = _jobs([("bold", "Acme Corp\tPune, India"), ("p", "Data Scientist | Aug 2022 - Present"), ("b", "Built x."),
                     ("p", "Data Analyst | Jan 2020 - Jul 2022"), ("b", "Built y.")], tmp_path)
    assert [(j[1], j[2]) for j in jobs] == [("Acme Corp", "Pune, India"), ("Acme Corp", "Pune, India")]


@pytest.mark.parametrize("title", ["Production Supervisor", "Delivery Driver", "Paralegal", "Barista",
                                   "Assistant Professor", "Postdoc", "Logistics NCOIC", "Family Caregiver"])
def test_role_words_outside_tech(title):
    assert ResumeNormalizer()._looks_like_title(title)


def test_a_title_with_a_comma_over_a_company_line_p916(tmp_path):
    """The ATS template's own output, read back: "SVP, Global Supply Chain<tab>dates"
    then the company alone (no location) is one title, not "Title, Company"."""
    jobs, _ = _jobs([("bold", "SVP, Global Supply Chain\t2013 – 2018"), ("p", "Meridian Consumer Products"),
                     ("b", "Led S&OP redesign across 4 sites."),
                     ("bold", "Shift Supervisor, Starbucks, Atlanta GA\t2021 – Present"),
                     ("p", "Opened and closed the store"), ("b", "Trained 6 baristas.")], tmp_path)
    assert jobs[0][:3] == ("SVP, Global Supply Chain", "Meridian Consumer Products", None)
    assert jobs[1][:3] == ("Shift Supervisor", "Starbucks", "Atlanta GA")  # a sentence is never the company
