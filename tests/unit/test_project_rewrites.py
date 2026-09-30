"""Project bullets go through the same plan -> rewrite -> validate flow (P1.7)."""
from unittest.mock import MagicMock

from app.analysis.rewriter import LLMRewriter
from app.analysis.tailor_planner import TailoringPlanner
from app.domain.evidence import Evidence
from app.domain.job import JobDescription, Requirement
from app.domain.resume import Candidate, Experience, Project, Resume, ResumeBullet
from app.llm.schemas import RoleBulletRewrite, RoleRewriteResult
from app.services.tailor import TailorService


def _setup():
    resume = Resume(
        candidate=Candidate(name="A"),
        experience=[Experience(id="e1", company="Acme", title="Analyst", bullets=[
            ResumeBullet(id="e1_b1", text="Built Python reports for sales.")])],
        projects=[Project(id="p1", name="Movie Recommender", bullets=[
            ResumeBullet(id="p1_b1", text="Trained a PyTorch ranking model on movie ratings."),
            ResumeBullet(id="p1_b2", text="Painted the office mural."),
        ])],
    )
    evidence = [Evidence(id="ev_e1", source_type="experience", source_id="e1_b1", text="Acme: Built Python reports for sales."),
                Evidence(id="ev_p1", source_type="project", source_id="p1_b1",
                         text="Project (Movie Recommender): Trained a PyTorch ranking model on movie ratings."),
                Evidence(id="ev_p2", source_type="project", source_id="p1_b2",
                         text="Project (Movie Recommender): Painted the office mural.")]
    req = "Build ranking models in PyTorch and Python"
    job = JobDescription(keywords=["PyTorch", "Python", "ranking"], requirements=[Requirement(id="r1", text=req)],
                         raw_text=req)
    return resume, evidence, job


def test_planner_scores_project_bullets():
    resume, evidence, job = _setup()
    actions = {a.source_id: a for a in TailoringPlanner().create_plan(resume, job, evidence, []).actions}
    assert actions["p1_b1"].action == "REWRITE" and actions["p1_b1"].target_section == "projects"
    assert actions["p1_b2"].action == "KEEP"
    assert "PyTorch" in actions["p1_b1"].keywords


def test_projects_get_one_extra_call_with_project_names():
    resume, evidence, job = _setup()
    plan = TailoringPlanner().create_plan(resume, job, evidence, [])
    llm = MagicMock()
    llm.is_available.return_value = True
    llm.generate_json.side_effect = [
        RoleRewriteResult(bullets=[RoleBulletRewrite(bullet_id="e1_b1", rewritten="Automated Python sales reports.")]),
        RoleRewriteResult(bullets=[RoleBulletRewrite(bullet_id="p1_b1",
                                                     rewritten="Trained a PyTorch ranking model on movie ratings data.")]),
    ]
    proposals = {p.target_semantic_id: p for p in LLMRewriter(llm).execute_plan(resume, plan, evidence, job)}
    assert llm.generate_json.call_count == 2  # one for the job, one for all projects
    project_prompt = llm.generate_json.call_args_list[1].kwargs["messages"][1]["content"]
    assert "Section: Projects" in project_prompt and "sub-heading: Movie Recommender" in project_prompt
    assert proposals["p1_b1"].status == "ok"


def test_tailor_applies_project_rewrite(tmp_path):
    from app.services.run_manager import RunManager
    resume, evidence, job = _setup()
    service = TailorService(llm_client=MagicMock(is_available=MagicMock(return_value=False)))
    service.run_manager = RunManager(base_runs_dir=str(tmp_path / "runs"))
    service.pdf_converter.convert_docx_to_pdf = MagicMock(return_value=None)
    from app.domain.resume_document import ResumeDocument
    parsed = (MagicMock(document_map=None), ResumeDocument(resume=resume), evidence)
    prop = {"target_semantic_id": "p1_b1", "target_source_location_id": "p1_b1", "original_text": "x",
            "proposed_text": "Trained a PyTorch ranking model on movie ratings data.", "evidence_ids": ["ev_p1"]}
    result = service.tailor_resume("x.pdf", "Requirements:\n- PyTorch.", str(tmp_path / "out"), mode="ATS_DEFAULT",
                                   preapproved_proposals=[prop], parsed=parsed, job_desc=job)
    import docx
    texts = [p.text for p in docx.Document(result["docx"]).paragraphs]
    assert "Trained a PyTorch ranking model on movie ratings data." in texts
