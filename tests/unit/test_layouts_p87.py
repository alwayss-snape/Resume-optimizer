"""P8.7: text boxes, a name in a separated page header, "·" bullets."""
import docx
from docx.oxml import parse_xml

from app.analysis.resume_normalizer import ResumeNormalizer
from app.ingestion.docx import DocxParser
from app.ingestion.pdf import PdfParser

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_BOX = ('<w:r ' + _W + ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>'
        '<wp:inline><wp:extent cx="1" cy="1"/><wp:docPr id="1" name="TB"/><a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:txbx><w:txbxContent>'
        '{body}</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline>'
        '</w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>{body}</w:txbxContent>'
        '</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>')


def _p(text, bold=False):
    rpr = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f"<w:p><w:r>{rpr}<w:t>{text}</w:t></w:r></w:p>"


def test_text_box_content_is_read_once_in_order(tmp_path):
    d = docx.Document()
    d.add_paragraph().add_run("Sam Taylor").bold = True
    d.add_paragraph("sam@example.com | 555-555-0100")
    body = _p("EXPERIENCE", bold=True) + _p("Office Manager, Bright Dental, 2018 - Present") + _p(
        "- Managed scheduling with Dentrix")
    d.add_paragraph()._p.append(parse_xml(_BOX.format(body=body)))
    d.add_paragraph().add_run("EDUCATION").bold = True
    d.add_paragraph("BA English, State University, 2016")
    path = str(tmp_path / "r.docx")
    d.save(path)
    raw = DocxParser().parse(path)
    texts = [b.text for b in raw.blocks]
    assert texts.count("Office Manager, Bright Dental, 2018 - Present") == 1  # the fallback copy is skipped
    assert texts.index("EXPERIENCE") < texts.index("EDUCATION")
    resume = ResumeNormalizer().normalize(raw)[0].resume
    assert [(e.title, e.company, len(e.bullets)) for e in resume.experience] == [("Office Manager", "Bright Dental", 1)]


def test_name_from_a_separated_page_header(tmp_path):
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Lena Park · Product Designer · lena@example.com · lenapark.design"
    d.add_paragraph().add_run("SKILLS").bold = True
    d.add_paragraph("Figma")
    path = str(tmp_path / "r.docx")
    d.save(path)
    c = ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume.candidate
    assert (c.name, c.headline, c.links) == ("Lena Park", "Product Designer", ["lenapark.design"])


def test_middle_dot_bullets():
    assert PdfParser()._split_bullet("· Achieved 132% of quota") == "Achieved 132% of quota"
    _t, start, end = ResumeNormalizer()._parse_title_and_dates("Senior AE, CloudMetrics\tJan 2021 · Present".split("\t")[1])
    assert (start, end) == ("Jan 2021", "Present")
