"""P10.6: the Academic CV template. Appointments are jobs, publications are
kept word for word, in order and numbered, and sections follow CV order."""
import re
from unittest.mock import MagicMock

import docx

from app.analysis.cv_mode import academic_section_order, experience_heading, move_appointments_from_education
from app.domain.resume import Candidate, Education, Experience, OtherSection, Resume, SectionLine
from app.domain.resume_document import ResumeDocument, ResumePresentation
from app.rendering.html_renderer import HtmlResumeRenderer
from app.rendering.template_renderer import TemplateRenderer
from app.services.tailor import TailorService

PERSONA = "data/eval/personas/academic"


def _paragraphs(path):
    return [(p.style.name, p.text) for p in docx.Document(path).paragraphs if p.text.strip()]


def test_the_academic_persona_in_cv_order_with_publications_verbatim(tmp_path):
    source = [t for _s, t in _paragraphs(f"{PERSONA}/resume.docx")]
    service = TailorService(llm_client=None, keep_run=False)
    service.pdf_converter.convert_docx_to_pdf = MagicMock(return_value=None)
    out = service.tailor_resume(f"{PERSONA}/resume.docx", open(f"{PERSONA}/jd.txt").read(), str(tmp_path))
    paragraphs = _paragraphs(out["docx"])
    headings = [t for _s, t in paragraphs if t.isupper() and len(t) < 40 and t != "PHD"]
    assert headings[:5] == ["RESEARCH INTERESTS", "EDUCATION", "ACADEMIC APPOINTMENTS", "GRANTS AND FUNDING",
                            "PUBLICATIONS"]
    assert headings[5:] == ["TEACHING", "INVITED TALKS", "SERVICE", "PROFESSIONAL MEMBERSHIPS"]
    pubs = [t for t in source if re.match(r"\d+\. Okafor", t)]
    assert len(pubs) == 45
    assert [t for _s, t in paragraphs if t in pubs] == pubs  # every one, word for word, in its order
    assert out["target_pages"] is None


def test_unnumbered_publications_are_numbered_by_the_template_not_rewritten(tmp_path):
    resume = Resume(candidate=Candidate(name="A"), other_sections=[OtherSection(
        id="p", heading="Publications", lines=[SectionLine(text="Okafor H. (2020). A paper.", bullet=True),
                                               SectionLine(text="Okafor H. (2021). Another.", bullet=True)])])
    path = str(tmp_path / "cv.docx")
    academic = ResumePresentation(cv_mode="academic", section_order=academic_section_order(resume))
    TemplateRenderer().render_ats_default(ResumeDocument(resume=resume, presentation=academic), path)
    assert ("List Number", "Okafor H. (2020). A paper.") in _paragraphs(path)
    assert "<ol><li>Okafor H. (2020). A paper.</li>" in HtmlResumeRenderer().render(
        ResumeDocument(resume=resume, presentation=academic))
    TemplateRenderer().render_ats_default(ResumeDocument(resume=resume), path)  # a standard resume: bullets
    assert ("List Bullet", "Okafor H. (2020). A paper.") in _paragraphs(path)


def test_an_appointment_read_as_education_becomes_a_job():
    resume = Resume(candidate=Candidate(name="A"), experience=[
        Experience(id="e1", company="State University", title="Assistant Professor", start_date="2021",
                   end_date="Present")],
        education=[Education(id="d1", institution="Harvard University", degree="Postdoctoral Fellow",
                             dates="2018 – 2021"),
                   Education(id="d2", institution="Cornell University", degree="Ph.D. Microbiology", dates="2018")])
    notes = move_appointments_from_education(resume)
    assert notes == ['Moved "Postdoctoral Fellow" from Education to your appointments (an Academic CV lists it as a job).']
    assert [(e.title, e.company, e.start_date, e.end_date) for e in resume.experience] == [
        ("Assistant Professor", "State University", "2021", "Present"),
        ("Postdoctoral Fellow", "Harvard University", "2018", "2021")]
    assert [e.degree for e in resume.education] == ["Ph.D. Microbiology"]
    assert experience_heading(resume) == "Academic Appointments"
    resume.experience.append(Experience(id="e3", company="Pivot AgBio", title="Senior Scientist"))
    assert experience_heading(resume) == "Professional Experience"
