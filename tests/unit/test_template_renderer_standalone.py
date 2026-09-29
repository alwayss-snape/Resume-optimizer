import os
import docx

from app.domain.resume import Candidate, Education, Experience, Resume, ResumeBullet
from app.domain.resume_document import ResumeDocument
from app.rendering.template_renderer import TemplateRenderer


def _full_text(path):
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs)


def test_template_renderer_ats_mode(tmp_path):
    resume = Resume(
        candidate=Candidate(name="John Smith", email="john@example.com"),
        summary="Experienced Software Engineer",
        experience=[
            Experience(
                id="exp_1",
                company="TechCorp",
                title="Lead Developer",
                start_date="2020",
                end_date="Present",
                bullets=[ResumeBullet(id="b1", text="Led cloud architecture on AWS.")],
            )
        ],
    )

    renderer = TemplateRenderer()
    out_path = str(tmp_path / "ats_resume.docx")
    res_path = renderer.render_ats_default(resume, out_path)

    assert os.path.exists(res_path)
    full_text = _full_text(res_path)
    assert "John Smith" in full_text
    assert "Lead Developer" in full_text
    assert "TechCorp" in full_text
    assert "2020" in full_text and "Present" in full_text
    assert "Led cloud architecture on AWS." in full_text


def test_template_renderer_accepts_resumedocument(tmp_path):
    resume = Resume(
        candidate=Candidate(name="Jane Doe", email="jane@example.com"),
        summary="Senior Engineer",
        experience=[
            Experience(
                id="exp_2",
                company="ExampleCo",
                title="SWE",
                bullets=[ResumeBullet(id="b2", text="Built scalable APIs.")],
            )
        ],
    )

    resume_doc = ResumeDocument(resume=resume)
    renderer = TemplateRenderer()
    out_path = str(tmp_path / "ats_resume_doc.docx")
    res_path = renderer.render_ats_default(resume_doc, out_path)

    assert os.path.exists(res_path)
    full_text = _full_text(res_path)
    assert "Jane Doe" in full_text
    assert "Built scalable APIs." in full_text


def test_template_renderer_handles_missing_optional_fields(tmp_path):
    """No dates, no location, no email — must not crash, and must still
    produce a readable document (regression guard for the new right-tab
    date layout, which skips the tab stop entirely when dates is empty)."""
    resume = Resume(
        candidate=Candidate(name="No Dates Person"),
        experience=[
            Experience(id="exp_1", company="Acme", title="Contractor", bullets=[]),
        ],
    )
    renderer = TemplateRenderer()
    out_path = str(tmp_path / "no_dates.docx")
    res_path = renderer.render_ats_default(resume, out_path)
    full_text = _full_text(res_path)
    assert "No Dates Person" in full_text
    assert "Contractor" in full_text
    assert "Acme" in full_text
