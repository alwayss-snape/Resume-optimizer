"""P2.5: the rendered template must read back exactly as rendered, and the
round-trip check must fail loudly when it doesn't."""
import shutil

import docx
import pytest

from app.domain.resume import Candidate, Education, Experience, Resume, ResumeBullet, Role
from app.domain.resume_document import ResumeDocument
from app.rendering.pdf_converter import PdfConverter
from app.rendering.template_renderer import TemplateRenderer
from app.validation.output import OutputQAValidator


def _resume():
    return Resume(
        candidate=Candidate(name="Avery Lee", email="avery@example.com", phone="+1 555 010 0100",
                            location="Pune, India", links=["https://www.linkedin.com/in/avery"]),
        summary="Engineer who builds data platforms.",
        experience=[
            Experience(id="e1", company="Acme Corp", title="Engineer II", location="Pune, India",
                       roles=[Role(title="Engineer II", start_date="Aug 2024", end_date="Present"),
                              Role(title="Engineer I", start_date="Jan 2022", end_date="Jul 2024")],
                       bullets=[ResumeBullet(id="b1", text="Built the ingestion service in Python.", group="Data Platform"),
                                ResumeBullet(id="b2", text="Cut pipeline cost by 30% with Spark tuning.", group="Data Platform"),
                                ResumeBullet(id="b3", text="Shipped a fraud model to production.", group="Fraud Detection")]),
            Experience(id="e2", company="Globex", title="Analyst", roles=[
                Role(title="Analyst", start_date="Jun 2020", end_date="Dec 2021")],
                bullets=[ResumeBullet(id="b4", text="Automated weekly reporting in SQL.")]),
        ],
        education=[Education(id="ed1", institution="State University", degree="B.Tech in Computer Science",
                             location="Chennai, India", dates="2016 - 2020")],
        skills={"Languages": ["Python", "SQL"], "Tools": ["Spark", "Docker"]},
        certifications=[{"name": "Cloud Practitioner"}],
    )


def _render(tmp_path, resume):
    return TemplateRenderer().render_ats_default(ResumeDocument(resume=resume), str(tmp_path / "r.docx"))


def test_template_docx_round_trips_cleanly(tmp_path):
    resume = _resume()
    assert OutputQAValidator().round_trip(_render(tmp_path, resume), resume) == []


@pytest.mark.skipif(not PdfConverter().find_libreoffice_binary(), reason="needs LibreOffice")
def test_template_pdf_round_trips_cleanly(tmp_path):
    resume = _resume()
    pdf = PdfConverter().convert_docx_to_pdf(_render(tmp_path, resume), str(tmp_path))
    assert OutputQAValidator().round_trip(pdf, resume) == []


def test_round_trip_flags_what_does_not_read_back(tmp_path):
    rendered = _render(tmp_path, _resume())
    expected = _resume()
    expected.candidate.email = "other@example.com"
    expected.experience[0].bullets.append(ResumeBullet(id="bx", text="A bullet that was never rendered."))
    expected.experience[1].roles[0].start_date = "Jan 2019"
    expected.skills["Tools"].append("Kubernetes")
    problems = OutputQAValidator().round_trip(rendered, expected)
    assert all(p.startswith("ATS round-trip (DOCX):") for p in problems)
    text = " | ".join(problems)
    assert "email not found" in text
    assert "1 of 5 bullets not read back" in text
    assert "role titles or dates differ" in text
    assert "Kubernetes" in text


def test_round_trip_flags_a_broken_file(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx")
    problems = OutputQAValidator().round_trip(str(path), _resume())
    assert problems and "could not be re-parsed" in problems[0]


def test_bold_label_skill_line_is_not_a_heading(tmp_path):
    """DOCX parser: 'Tools: Spark, Docker' with only the label bold is content."""
    from app.ingestion.docx import DocxParser
    d = docx.Document()
    d.add_paragraph().add_run("SKILLS").bold = True
    p = d.add_paragraph()
    p.add_run("Tools: ").bold = True
    p.add_run("Spark, Docker")
    path = str(tmp_path / "s.docx")
    d.save(path)
    blocks = DocxParser().parse(path).blocks
    tools = next(b for b in blocks if b.text.startswith("Tools"))
    assert tools.block_type != "heading"


def test_normalizer_reads_title_first_layout():
    """Title<tab>dates, then 'Company · Location'; a second job's company line
    replaces the one carried over from the previous job."""
    from app.analysis.resume_normalizer import ResumeNormalizer
    from app.ingestion.docx import DocxParser
    import tempfile, os
    d = docx.Document()
    d.add_paragraph().add_run("Avery Lee").bold = True
    d.add_paragraph().add_run("WORK EXPERIENCE").bold = True
    for title, dates, company in [("Engineer II", "Aug 2024 – Present", "Acme Corp · Pune, India"),
                                  ("Analyst", "Jun 2020 – Dec 2021", "Globex")]:
        t = d.add_paragraph()
        t.add_run(title).bold = True
        t.add_run("\t" + dates)
        d.add_paragraph(company)
        d.add_paragraph("Did the work well.", style="List Bullet")
    d.add_paragraph().add_run("EDUCATION").bold = True
    e = d.add_paragraph()
    e.add_run("B.Tech in Computer Science").bold = True
    e.add_run("\t2016 – 2020")
    d.add_paragraph("State University · Chennai, India")
    path = os.path.join(tempfile.mkdtemp(), "t.docx")
    d.save(path)
    resume = ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume
    jobs = [(e.company, e.location, e.title, e.start_date, e.end_date) for e in resume.experience]
    assert jobs == [("Acme Corp", "Pune, India", "Engineer II", "Aug 2024", "Present"),
                    ("Globex", None, "Analyst", "Jun 2020", "Dec 2021")]
    edu = resume.education[0]
    assert (edu.degree, edu.institution, edu.location, edu.dates) == (
        "B.Tech in Computer Science", "State University", "Chennai, India", "2016 – 2020")


def test_degree_dates_are_not_also_the_location(tmp_path):
    """Live judge finding: institution line, then 'Degree<tab>dates' put the
    dates into location too, so the output printed them twice."""
    from app.analysis.resume_normalizer import ResumeNormalizer
    from app.ingestion.docx import DocxParser
    d = docx.Document()
    d.add_paragraph().add_run("Avery Lee").bold = True
    d.add_paragraph().add_run("EDUCATION").bold = True
    d.add_paragraph().add_run("Lakeside State University").bold = True
    d.add_paragraph("B.S. in Computer Science\t2021 - 2025")
    path = str(tmp_path / "e.docx")
    d.save(path)
    edu = ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume.education[0]
    assert (edu.institution, edu.degree, edu.dates, edu.location) == (
        "Lakeside State University", "B.S. in Computer Science", "2021 - 2025", None)


def test_institution_dates_are_not_the_location(tmp_path):
    from app.analysis.resume_normalizer import ResumeNormalizer
    from app.ingestion.docx import DocxParser
    d = docx.Document()
    d.add_paragraph().add_run("Avery Lee").bold = True
    d.add_paragraph().add_run("EDUCATION").bold = True
    d.add_paragraph("Lakeside State University\t2021 - 2025")
    d.add_paragraph("B.S. in Computer Science\tChennai, India")
    path = str(tmp_path / "e2.docx")
    d.save(path)
    edu = ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume.education[0]
    assert (edu.institution, edu.degree, edu.dates, edu.location) == (
        "Lakeside State University", "B.S. in Computer Science", "2021 - 2025", "Chennai, India")
