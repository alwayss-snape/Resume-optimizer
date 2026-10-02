"""P8.8: one-line education entries stay separate, details stay with their
entry, and degrees are listed newest first."""
import docx
import pytest

from app.analysis.resume_normalizer import ResumeNormalizer
from app.ingestion.docx import DocxParser


def _education(lines, tmp_path):
    d = docx.Document()
    d.add_paragraph().add_run("Avery Lee").bold = True
    d.add_paragraph().add_run("EDUCATION").bold = True
    for line in lines:
        d.add_paragraph(line)
    path = str(tmp_path / "r.docx")
    d.save(path)
    return ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume.education


@pytest.mark.parametrize("line, want", [
    ("J.D., cum laude, Georgetown University Law Center, 2017", ("J.D., cum laude", "Georgetown University Law Center", "2017")),
    ("MBA, Kellogg School of Management, Northwestern University, 2005",
     ("MBA", "Kellogg School of Management, Northwestern University", "2005")),
    ("B.E. Computer Science, Anna University, 2015 — 78%", ("B.E. Computer Science, 78%", "Anna University", "2015")),
    ("High School Diploma, Whitehaven High School, 2017", ("High School Diploma", "Whitehaven High School", "2017")),
    ("BFA Graphic Design, Rhode Island School of Design, 2017",
     ("BFA Graphic Design", "Rhode Island School of Design", "2017")),
    # The part after the degree must name a school; otherwise the school is on another line.
    ("Bachelor of Science, Electrical Engineering, May 2020", None),
    ("M.S., Computer Science, 2022", None),
    ("Grado en Enfermería, Universidad Complutense de Madrid, 2019",
     ("Grado en Enfermería", "Universidad Complutense de Madrid", "2019")),
])
def test_one_line_entries(line, want):
    assert ResumeNormalizer()._split_education_line(line) == want


def test_two_degrees_stay_two_entries_newest_first(tmp_path):
    edu = _education(["B.S. Biology, University of Texas at Austin, 2016",
                      "M.Ed. Curriculum and Instruction, Texas State University, 2019"], tmp_path)
    assert [(e.degree, e.institution, e.dates) for e in edu] == [
        ("M.Ed. Curriculum and Instruction", "Texas State University", "2019"),
        ("B.S. Biology", "University of Texas at Austin", "2016")]


def test_coursework_is_a_detail_and_semicolons_split_entries(tmp_path):
    edu = _education(["B.S. Statistics, Boston University, Expected May 2026, GPA 3.7/4.0",
                      "Relevant Coursework: Regression Analysis, Data Mining"], tmp_path)
    assert len(edu) == 1 and edu[0].details == ["Relevant Coursework: Regression Analysis, Data Mining"]
    assert edu[0].dates == "Expected May 2026"
    edu = _education(["Advanced Leader Course (ALC), 2019; B.S. Business Administration (in progress), "
                      "University of Maryland Global Campus"], tmp_path)
    assert [e.degree or e.institution for e in edu] == ["Advanced Leader Course (ALC)",
                                                       "B.S. Business Administration (in progress)"]


def test_school_line_then_degree_line_pairs_up(tmp_path):
    edu = _education(["Purdue University, West Lafayette, IN", "Bachelor of Science, Electrical Engineering, May 2020",
                      "University of Illinois Urbana-Champaign", "M.S., Computer Science, 2022"], tmp_path)
    assert [(e.institution, e.degree, e.dates) for e in edu] == [
        ("University of Illinois Urbana-Champaign", "M.S., Computer Science", "2022"),
        ("Purdue University, West Lafayette, IN", "Bachelor of Science, Electrical Engineering", "May 2020")]
