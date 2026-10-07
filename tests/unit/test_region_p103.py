"""P10.3: the region comes from the JD's own words, with its evidence; it
sets the paper and date style; the user can change it in Arrange with no
LLM call; personal details only get advice."""
import docx
import pytest

from app.analysis.region import apply_region, personal_details_advice, suggest_region
from app.domain.resume import Candidate, Education, Experience, OtherSection, Resume, ResumeBullet, SectionLine
from app.domain.resume_document import ResumeDocument, ResumePresentation
from app.rendering.html_renderer import HtmlResumeRenderer
from app.rendering.layout import date_range, format_date_text
from app.rendering.template_renderer import TemplateRenderer
from app.validation.content_lint import lint


@pytest.mark.parametrize("jd, region, evidence", [
    ("Senior Analyst\nLocation: Austin, TX 78701\nSalary: $120,000", "us", "Austin, TX 78701"),
    ("Apply on USAJOBS. Series GS-0343-12.", "us", "USAJOBS"),
    ("Hybrid role based in Denver, Colorado.", "us", "Denver, Colorado"),
    ("Remote - US only. Python, R and SQL.", "us", "Remote - US"),
    ("Data Engineer, Bengaluru. CTC 18 LPA.", "india", "Bengaluru"),
    ("Pune, IN (hybrid)", "india", "Pune"),
    ("Backend developer in London, hybrid, £60,000", "uk_eu", "London"),
    ("Wir suchen ... Standort: Berlin", "uk_eu", "Berlin"),
    ("Salary €55k, fully remote within the EU", "uk_eu", "EU"),
    ("Strong Python, SQL, OR tools and IN-house ML.", "other", None),
    ("We build great software. Join us!", "other", None),
])
def test_region_from_the_jd_with_its_evidence(jd, region, evidence):
    guess = suggest_region(jd)
    assert (guess.region, guess.evidence) == (region, evidence)


def test_apply_region_sets_paper_and_dates():
    us = apply_region(ResumePresentation(), "us")
    assert (us.page_size, us.date_style, us.region) == ("Letter", "numeric", "us")
    for region in ("uk_eu", "india", "other", None, "nonsense"):
        p = apply_region(ResumePresentation(), region)
        assert (p.page_size, p.date_style) == ("A4", "month")


def test_numeric_dates_keep_years_and_present():
    assert date_range("January 2022", "Current") == "Jan 2022 – Present"
    assert date_range("January 2022", "Current", "numeric") == "01/2022 – Present"
    assert date_range("Aug 2016", "May 2020", "numeric") == "08/2016 – 05/2020"
    assert format_date_text("2016 - 2020", "numeric") == "2016 – 2020"
    assert format_date_text("Expected May 2026", "numeric") == "Expected 05/2026"


def _resume():
    return Resume(
        candidate=Candidate(name="A Person", details=["Date of Birth: 01/01/1990", "Father's Name: B Person"]),
        experience=[Experience(id="e1", company="Acme", title="Engineer", start_date="Jan 2022",
                               end_date="Present", bullets=[ResumeBullet(id="b1", text="Built a billing service.")])],
        education=[Education(id="ed1", institution="State University", degree="BSc", dates="2016 - 2020")],
        other_sections=[OtherSection(id="o1", heading="Declaration", lines=[
            SectionLine(text="I hereby declare that the above is true.")])],
    )


def test_us_resume_renders_on_letter_with_numeric_dates(tmp_path):
    doc = ResumeDocument(resume=_resume(), presentation=apply_region(ResumePresentation(), "us"))
    path = str(tmp_path / "us.docx")
    TemplateRenderer().render_ats_default(doc, path)
    text = "\n".join(p.text for p in docx.Document(path).paragraphs)
    assert "01/2022 – Present" in text and "Jan 2022" not in text
    html = HtmlResumeRenderer().render(doc)
    assert "01/2022 – Present" in html and "size: letter;" in html
    # The default stays as it was.
    TemplateRenderer().render_ats_default(ResumeDocument(resume=_resume()), path)
    assert "Jan 2022 – Present" in "\n".join(p.text for p in docx.Document(path).paragraphs)


def test_personal_details_get_advice_only_for_us_and_europe():
    resume = _resume()
    advice = personal_details_advice(resume, "us")
    assert "date of birth" in advice and "father's or spouse's name" in advice and "a declaration" in advice
    assert "US employers" in advice
    assert personal_details_advice(resume, "india") is None
    assert personal_details_advice(resume, None) is None
    checks = [i.check for i in lint(resume, region="uk_eu").issues]
    assert "personal_details" in checks
    assert "personal_details" not in [i.check for i in lint(resume).issues]
    # Advice only: nothing removed.
    assert len(resume.candidate.details) == 2 and resume.other_sections


def test_the_api_suggests_the_region_and_arrange_switches_it(tmp_path):
    """Drafting returns the suggestion with its evidence; tailoring uses the
    region the user confirmed; Arrange changes it with no LLM call."""
    from fastapi.testclient import TestClient
    from unittest.mock import MagicMock

    from app.api.main import create_app
    from app.services.tailor import TailorService
    from app.services.run_manager import RunManager
    from tests.unit.test_api import SAMPLE_DOCX, _events, _upload

    def make_service(model=None):
        service = TailorService(llm_client=None)
        service.run_manager = RunManager(base_runs_dir=str(tmp_path / "runs"))
        service.pdf_converter.convert_docx_to_pdf = MagicMock(return_value=None)
        return service

    jd = ("Senior Python Developer\nLocation: Austin, TX\nWe need Python, Django, PostgreSQL, REST APIs, Docker "
          "and AWS experience. 5+ years building backend services.")
    with TestClient(create_app(make_service=make_service, serve_web=False)) as client:
        assert client.post("/api/parse", files=_upload(SAMPLE_DOCX), data={"jd_text": jd}).status_code == 200
        drafted = _events(client.post("/api/proposals", json={}))[-1][1]
        assert drafted["region"] == {"region": "us", "label": "US (Letter)", "evidence": "Austin, TX"}
        assert client.post("/api/tailor", json={"region": "mars"}).status_code == 422
        final = _events(client.post("/api/tailor", json={"region": "us"}))[-1][1]
        assert final["region"] == {"region": "us", "label": "US (Letter)", "evidence": "Austin, TX"}
        layout = final["arrangement"]["layout"]
        assert layout["region"] == "us"
        out = client.post("/api/arrange", json={"layout": {**layout, "region": "india"}}).json()
        assert out["region"] == {"region": "india", "label": "India (A4)", "evidence": None}
        assert client.post("/api/arrange", json={"layout": {**layout, "region": "mars"}}).status_code == 422
