"""P10.13: choose projects, not lines."""
from unittest.mock import MagicMock

from fastapi.testclient import TestClient

from app.analysis.project_select import (impact, jd_is_thin, keyword_fit, left_out_ids, remove_bullets,
                                         select_projects)
from app.domain.job import JobDescription
from app.domain.resume import Candidate, Experience, Resume, ResumeBullet
from app.domain.tailoring import TailoringAction
from app.services.tailor import TailorService

REPLICA = "tests/fixtures/resumes/replica_layout.pdf"
JD = "Senior Data Scientist. Pricing, yield optimization and inventory management at scale."
# A current job pasted on Check details as notes under four project headings.
BANK = {"company": "Northwind Grocers", "title": "ML Engineer", "location": "Bengaluru", "current": True,
        "start_date": "March 2026", "end_date": "Present", "description": (
            "Label Compliance:\nDeployed an OCR pipeline validating 19+ compliance fields.\n"
            "Customer Decisioning:\nBuilt production scoring pipeline for 12M users replacing 3 packages.\n"
            "Engineered 140 features from 2.3B transaction rows.\n"
            "Demand Forecasting:\nPiloting a global demand forecasting model on 100-200 model IDs.\n"
            "Own Brands:\nReconciled database and shelf brand names with LLM matching.\n")}


def _job(exp_id, groups, end="Present"):
    bullets = [ResumeBullet(id=f"{exp_id}_{(g or 'x')[:3]}{n}", text=t, group=g)
               for g, texts in groups for n, t in enumerate(texts)]
    return Experience(id=exp_id, company=exp_id.title(), title="Engineer", start_date="Jan 2020", end_date=end,
                      bullets=bullets)


def _actions(resume, relevance):
    return [TailoringAction(action="REWRITE", source_id=b.id, relevance=relevance.get(b.group, 0.0))
            for e in resume.experience for b in e.bullets]


def _jd(n_keywords):
    return JobDescription(raw_text="x", job_title="Data Scientist",
                          keywords=["Data", "Data Scientist", *[f"skill {i}" for i in range(n_keywords)]])


def test_a_name_is_not_early_stage_but_piloting_is():
    deal, signals = impact(["Route Pilot AI Copilot",
                            "Won an internal hackathon with Route Pilot, an AI procurement copilot."])
    assert deal > 0.3 and "an award or win" in signals and not any("early" in s for s in signals)
    pilot, signals = impact(["Demand Forecasting", "Piloting a global model on 100-200 model IDs."])
    assert pilot == 0.0 and any("early stage" in s for s in signals)


def test_production_needs_shipped_work_not_the_word():
    assert impact(["Monitor", "Built agents that monitor production ML pipelines."])[0] == 0.0
    shipped, signals = impact(["Scoring", "Built production scoring pipeline replacing 3 legacy packages."])
    assert shipped >= 0.35 and "in production" in signals


def test_numbers_count_once_and_years_are_not_scale():
    one = impact(["X", "Engineered features from 2.3B rows for 12M users."])[0]
    many = impact(["X", "Rows: 2.3B rows, 12M users, 140 features, 10,000 jobs, 60% of 90% cases."])[0]
    assert one == many  # the biggest list of figures doesn't win by itself
    assert impact(["X", "Worked on it from 2022 to 2024."])[0] == 0.0


def test_a_jd_naming_few_skills_ranks_by_impact():
    # "Data" and the job title don't count; a short post naming three skills is enough to rank by.
    assert jd_is_thin(_jd(2)) and not jd_is_thin(_jd(3))


def test_keyword_fit_compares_one_skill_at_a_time():
    vec = {"inventory management": [1, 0, 0], "pricing": [0, 1, 0], "Forecast": [0.9, 0.1, 0], "Demand plan": [0.8, 0, 0.2],
           "Labels": [0, 0, 1], "OCR fields": [0.1, 0, 0.9]}
    embed = lambda texts: [vec[t] for t in texts]
    (near, kws), (far, _) = keyword_fit([["Forecast", "Demand plan"], ["Labels", "OCR fields"]],
                                        ["inventory management", "pricing"], embed)
    assert near > far and kws[0] == "inventory management"
    assert keyword_fit([["a"]], ["pricing"], None) == [(0.0, [])]


def test_the_current_job_keeps_three_projects_and_an_older_one_two():
    resume = Resume(candidate=Candidate(name="A B"), experience=[
        _job("now", [("Alpha", ["a"]), ("Beta", ["b"]), ("Gamma", ["c"]), ("Delta", ["d"])]),
        _job("old", [("Kilo", ["k"]), ("Lima", ["l"]), ("Mike", ["m"])], end="Mar 2020"),
        _job("flat", [(None, ["x", "y", "z", "w"])], end="Jan 2019"),  # no projects: nothing to choose
    ])
    relevance = {"Alpha": 0.9, "Beta": 0.1, "Gamma": 0.5, "Delta": 0.7, "Kilo": 0.2, "Lima": 0.8, "Mike": 0.6}
    choices = select_projects(resume, _actions(resume, relevance), _jd(8))
    kept = {c.name for c in choices if c.chosen}
    assert kept == {"Alpha", "Gamma", "Delta", "Lima", "Mike"}
    assert [c.name for c in choices] == ["Alpha", "Beta", "Gamma", "Delta", "Kilo", "Lima", "Mike"]  # file order
    assert next(c for c in choices if c.name == "Beta").reason == "Less relevant to this job than the projects kept"
    assert left_out_ids(choices, None) == ["now_Bet0", "old_Kil0"]
    assert left_out_ids(choices, ["now::Alpha"]) == ["now_Alp0"]  # the user's choice wins


def test_removing_a_project_takes_its_heading_along():
    resume = Resume(candidate=Candidate(name="A B"), experience=[_job("now", [("Alpha", ["a1", "a2"]), ("Beta", ["b"])])])
    removed = remove_bullets(resume, ["now_Alp0", "now_Alp1"])
    assert removed == ["a1", "a2", "Alpha"] and [b.text for b in resume.experience[0].bullets] == ["b"]


def test_an_added_job_reads_project_headings():
    lines = TailorService._project_lines("Forecasting:\nBuilt a model.\nTuned it.\n\n# Pricing\nSet prices.\n")
    assert lines == [("Forecasting", "Built a model."), ("Forecasting", "Tuned it."), ("Pricing", "Set prices.")]
    assert TailorService._project_lines("Did A. Then did B.") == [(None, "Did A."), (None, "Then did B.")]


def test_an_added_project_bank_leaves_out_a_project_and_loses_nothing(tmp_path):
    service = TailorService(llm_client=None)
    jd = JD
    parsed, _ = service.apply_parse_corrections(service.parse_resume(REPLICA), {"added_jobs": [BANK]})
    added = parsed[1].resume.experience[0]
    assert added.company == "Northwind Grocers" and len([g for g, _ in added.bullet_groups()]) == 4
    result = service.tailor_resume(REPLICA, jd, str(tmp_path), preapproved_proposals=[], parsed=parsed,
                                   remember_answers=False)
    assert result["coverage"]["lost"] == [] and result["success"] is True
    assert any("Left out 1 project(s)" in line for line in open(result["changes_md"], encoding="utf-8"))
    # Brought back by the user: all four stay.
    kept_all = service.tailor_resume(REPLICA, jd, str(tmp_path / "all"), preapproved_proposals=[], parsed=parsed,
                                     remember_answers=False, left_out_projects=[])
    assert kept_all["coverage"]["lost"] == []
    assert not any("Left out" in line for line in open(kept_all["changes_md"], encoding="utf-8"))


def test_the_api_lists_projects_and_takes_the_users_choice(tmp_path):
    from app.api.main import create_app
    from app.services.run_manager import RunManager
    from tests.unit.test_api import _events

    def make_service(model=None):
        service = TailorService(llm_client=None)
        service.run_manager = RunManager(base_runs_dir=str(tmp_path / "runs"))
        service.pdf_converter.convert_docx_to_pdf = MagicMock(return_value=None)
        return service

    jd = JD
    with TestClient(create_app(make_service=make_service, serve_web=False)) as client, open(REPLICA, "rb") as f:
        assert client.post("/api/parse", files={"file": ("cv.pdf", f, "application/pdf")},
                           data={"jd_text": jd}).status_code == 200
        drafted = _events(client.post("/api/proposals", json={"corrections": {"added_jobs": [BANK]}}))[-1][1]
        projects = drafted["projects"]
        assert len(projects) == 4 and sum(not p["chosen"] for p in projects) == 1
        assert drafted["projects_ranked_by_impact"] is True
        assert all(p["reason"] for p in projects)
        dropped = next(p for p in projects if not p["chosen"])
        assert not any(p["target"] in dropped["bullet_ids"] for p in drafted["proposals"])  # not sent to the AI
        everything = client.post("/api/match-preview", json={"selection": [], "left_out": []}).json()
        suggested = client.post("/api/match-preview", json={"selection": []}).json()
        assert everything["rate"] >= suggested["rate"]
        final = _events(client.post("/api/tailor", json={"left_out": []}))[-1][1]
        assert final["coverage"]["lost"] == []


def test_page_fit_shortens_a_kept_project_but_never_removes_it():
    from app.rendering.page_fit import PageFitter
    resume = Resume(candidate=Candidate(name="A B"), experience=[
        _job("now", [("Alpha", ["a1", "a2", "a3"]), ("Beta", ["b1", "b2"]), ("Gamma", ["c1", "c2"])])])
    relevance = {b.id: 0.0 for b in resume.experience[0].bullets}
    kept = {"now::Beta", "now::Gamma"}
    saved, _ = PageFitter._trim_bullets(resume, 10_000, relevance, set(), kept_projects=kept)
    groups = [b.group for b in resume.experience[0].bullets]
    assert groups.count("Beta") == 2 and groups.count("Gamma") == 2  # kept whole: already at the minimum
    saved, notes = PageFitter._drop_sections(resume, 10_000, relevance, kept_projects=kept)
    assert {b.group for b in resume.experience[0].bullets} >= {"Beta", "Gamma"}


def test_a_heading_may_only_reword_the_projects_own_text():
    from app.analysis.project_select import heading_problem
    src = ["CDX Customer Decisioning Pipeline (12M Users)",
           "Built production scoring pipeline for 8 customer actions combining LightGBM propensity models"]
    assert heading_problem("Customer Propensity Scoring Pipeline", src) is None
    assert "Churn" in heading_problem("Customer Churn Prediction", src)
    assert heading_problem("Customer Pipeline for 12M Users", src) == "figures belong in the bullets"
    assert heading_problem("Customer Scoring Pipeline for Propensity Models At Scale", src) == "over 6 words"


def test_achievements_in_the_notes_go_to_achievements_not_a_project():
    service = TailorService(llm_client=None)
    bank = {**BANK, "description": BANK["description"] + "Achievements:\nWon an internal hackathon with Route Pilot.\n"}
    parsed, _ = service.apply_parse_corrections(service.parse_resume(REPLICA), {"added_jobs": [bank]})
    resume = parsed[1].resume
    assert resume.achievements[-1] == "Won an internal hackathon with Route Pilot."
    assert "Achievements" not in {b.group for b in resume.experience[0].bullets}


class _HeadingLLM:
    """Renames headings; rewrites nothing else."""
    provider, model, last_error, on_wait = "fake", "fake", None, None

    def is_available(self):
        return True

    def get_usage_summary(self):
        return {}

    def generate_json(self, messages, schema_model, **kwargs):
        from app.llm.schemas import HeadingRenameResult
        if schema_model is HeadingRenameResult:
            names = [l[len("- heading: "):] for l in messages[-1]["content"].splitlines() if l.startswith("- heading: ")]
            new = {"Customer Decisioning": "Customer Scoring Pipeline", "Label Compliance": "Label Compliance Churn Model"}
            return HeadingRenameResult(headings=[{"original": n, "renamed": new.get(n, n)} for n in names])
        raise RuntimeError("not needed in this test")


def test_an_accepted_heading_renames_the_project_and_a_made_up_one_is_dropped(tmp_path):
    service = TailorService(llm_client=_HeadingLLM())
    service.pdf_converter.convert_docx_to_pdf = MagicMock(return_value=None)
    parsed, _ = service.apply_parse_corrections(service.parse_resume(REPLICA), {"added_jobs": [BANK]})
    drafted = service.generate_proposals(REPLICA, JD, parsed=parsed)
    headings = {p.original_text: p for p in drafted["proposals"] if p.kind == "heading"}
    assert headings["Customer Decisioning"].proposed_text == "Customer Scoring Pipeline"
    assert headings["Customer Decisioning"].validation == "PASS"
    assert headings["Label Compliance"].validation == "REJECT"  # "Churn" and "Model" aren't in its text
    result = service.tailor_resume(REPLICA, JD, str(tmp_path), parsed=parsed, remember_answers=False,
                                   preapproved_proposals=[p.model_dump() for p in headings.values()])
    groups = {b.group for e in result["arrange"]["full_doc"].resume.experience for b in e.bullets}
    assert "Customer Scoring Pipeline" in groups and "Customer Decisioning" not in groups
    assert "Label Compliance Churn Model" not in groups


def test_each_new_job_starts_after_a_clear_gap():
    from docx import Document
    from app.rendering.template_renderer import TemplateRenderer
    resume = Resume(candidate=Candidate(name="A B"), experience=[
        _job("now", [("Alpha", ["a1"])]), _job("old", [("Kilo", ["k1"])], end="Mar 2020")])
    import tempfile, os
    path = os.path.join(tempfile.mkdtemp(), "r.docx")
    TemplateRenderer().render_ats_default(resume, path)
    titles = [p for p in Document(path).paragraphs if p.text.startswith("Engineer")]
    assert titles[1].paragraph_format.space_before.pt >= 12
    group = next(p for p in Document(path).paragraphs if p.text == "Kilo")
    assert group.runs[0].italic and group.runs[0].font.size.pt == 10


def test_a_project_left_out_on_review_can_be_brought_back_in_arrange(tmp_path):
    """P11.11: Arrange had no copy of left-out projects."""
    service = TailorService(llm_client=None)
    service.pdf_converter.convert_docx_to_pdf = MagicMock(return_value=None)
    parsed, _ = service.apply_parse_corrections(service.parse_resume(REPLICA), {"added_jobs": [BANK]})
    result = service.tailor_resume(REPLICA, JD, str(tmp_path), preapproved_proposals=[], parsed=parsed,
                                   remember_answers=False)
    state = result["arrange"]
    layout = state["default_layout"]
    out = layout.removed_bullets
    assert out  # the left-out project's bullets, ready to bring back
    full_texts = {b.id: b.text for e in state["full_doc"].resume.experience for b in e.bullets}
    assert all(i in full_texts for i in out)
    back = layout.model_copy(update={"removed_bullets": [], "trim": False})
    arranged = service.arrange(state, back, str(tmp_path / "arranged"))
    assert arranged["coverage"]["lost"] == []
    import html
    page = html.unescape(open(arranged["html"], encoding="utf-8").read())
    assert all(full_texts[i] in page for i in out)


def test_drafting_a_project_that_isnt_in_the_run_is_refused(tmp_path):
    from app.api.main import create_app
    from app.services.run_manager import RunManager
    from tests.unit.test_api import _events

    def make_service(model=None):
        service = TailorService(llm_client=None)
        service.run_manager = RunManager(base_runs_dir=str(tmp_path / "runs"))
        service.pdf_converter.convert_docx_to_pdf = MagicMock(return_value=None)
        return service

    with TestClient(create_app(make_service=make_service, serve_web=False)) as client, open(REPLICA, "rb") as f:
        client.post("/api/parse", files={"file": ("cv.pdf", f, "application/pdf")}, data={"jd_text": JD})
        drafted = _events(client.post("/api/proposals", json={"corrections": {"added_jobs": [BANK]}}))[-1][1]
        assert client.post("/api/draft-project", json={"key": "nope::Nothing"}).status_code == 404
        left_out = next(p for p in drafted["projects"] if not p["chosen"])
        assert client.post("/api/draft-project", json={"key": left_out["key"]}).json() == {"proposals": []}  # no AI here
