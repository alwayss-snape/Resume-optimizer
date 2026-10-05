"""P2.1: the ATS template (A4, Arial, standard headings, section order,
dates, contact line, skills cap, file name), shared by DOCX and HTML."""
from datetime import date

import docx
from docx.shared import Inches

from app.domain.resume import Candidate, Education, Experience, Project, Resume, ResumeBullet, Role
from app.domain.resume_document import ResumeDocument
from app.rendering.html_renderer import HtmlResumeRenderer
from app.rendering.layout import (
    contact_parts, date_range, display_skills, format_date, format_date_text, output_basename, section_order_for,
)
from app.rendering.template_renderer import TemplateRenderer

TODAY = date(2026, 9, 30)


def _resume(start="Jan 2019", end="Present"):
    return Resume(
        candidate=Candidate(name="Avery Lee", email="avery@example.com", phone="+1 555 0100",
                            location="Pune, India",
                            links=["https://github.com/avery", "https://www.linkedin.com/in/avery/"]),
        summary="Engineer who builds data platforms.",
        experience=[Experience(
            id="e1", company="Acme Corp", title="Engineer II", location="Pune, India",
            roles=[Role(title="Engineer II", start_date="August 2024", end_date="Current"),
                   Role(title="Engineer I", start_date=start, end_date="Jul 2024")],
            bullets=[ResumeBullet(id="b1", text="Built the ingestion service.")],
        )],
        projects=[Project(id="p1", name="Route Planner", bullets=[ResumeBullet(id="pb1", text="Wrote a solver.")])],
        education=[Education(id="ed1", institution="State University", degree="B.Tech", dates="2015 - 2019")],
        skills={"Languages": ["Python"], "Tools": ["Docker"]},
        certifications=[{"name": "Cloud Practitioner"}],
        interests=["Chess"],
    )


def test_format_date():
    assert format_date("August 2024") == "Aug 2024"
    assert format_date("Sept. 2020") == "Sep 2020"
    assert format_date("08/2021") == "Aug 2021"
    assert format_date("Current") == "Present"
    assert format_date("2019") == "2019"
    assert format_date("Summer 2020") == "Summer 2020"  # unrecognised wording is kept
    assert date_range("January 2022", "present") == "Jan 2022 – Present"
    assert date_range("2020", None) == "2020"
    assert format_date_text("2015 - 2019") == "2015 – 2019"
    assert format_date_text("Aug 2016 to May 2020") == "Aug 2016 – May 2020"


def test_section_order_default_and_early_career():
    order = section_order_for(_resume(), TODAY)
    assert order == ["summary", "experience", "skills", "education", "projects", "certifications",
                     "achievements", "interests"]
    junior = _resume()
    junior.experience[0].roles = [Role(title="Engineer I", start_date="Jan 2026", end_date="Present")]
    assert section_order_for(junior, TODAY).index("education") < section_order_for(junior, TODAY).index("experience")
    no_jobs = _resume()
    no_jobs.experience = []
    assert section_order_for(no_jobs, TODAY)[1] == "education"
    undated = _resume()
    undated.experience[0].roles = [Role(title="Engineer")]
    assert section_order_for(undated, TODAY)[1] == "experience"  # unknown length: keep the default


def test_contact_line_order():
    assert contact_parts(_resume().candidate) == [
        "avery@example.com", "+1 555 0100", "Pune, India", "linkedin.com/in/avery", "github.com/avery",
    ]


def test_skills_capped_at_four_lines():
    skills = {"A": ["a1"], "B": ["b1"], "C": ["c1"], "D": ["d1", "x"], "E": ["e1", "x"]}
    assert display_skills(skills) == {"A": ["a1"], "B": ["b1"], "C": ["c1"], "Other": ["d1", "x", "e1"]}
    assert display_skills({"A": ["a1"], "B": []}) == {"A": ["a1"]}


def test_output_basename():
    assert output_basename(_resume(), "Globex Inc.") == "Avery_Lee_Resume_Globex_Inc"
    assert output_basename(_resume(), None) == "Avery_Lee_Resume"
    assert output_basename(Resume(candidate=Candidate()), "") == "Resume"


def test_docx_template_page_font_order_and_dates(tmp_path):
    document = ResumeDocument(resume=_resume())
    out = TemplateRenderer().render_ats_default(document, str(tmp_path / "r.docx"))
    d = docx.Document(out)
    section = d.sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)
    assert section.left_margin == Inches(0.7) and section.top_margin == Inches(0.6)
    assert d.styles["Normal"].font.name == "Arial"
    assert not any(s.header.paragraphs[0].text for s in d.sections)  # no Word header
    assert not d.tables

    texts = [p.text for p in d.paragraphs if p.text]
    headings = [t for t in texts if t.isupper() and len(t) < 30]
    assert headings == ["SUMMARY", "WORK EXPERIENCE", "SKILLS", "EDUCATION", "PROJECTS", "CERTIFICATIONS", "INTERESTS"]
    assert texts[1] == "avery@example.com | +1 555 0100 | Pune, India | linkedin.com/in/avery | github.com/avery"
    assert "Engineer II\tAug 2024 – Present" in texts
    assert "Engineer I\tJan 2019 – Jul 2024" in texts
    assert "B.Tech\t2015 – 2019" in texts
    assert "State University" in texts


def test_docx_follows_presentation_order(tmp_path):
    document = ResumeDocument(resume=_resume())
    document.presentation.section_order = ["education", "experience", "summary"]
    out = TemplateRenderer().render_ats_default(document, str(tmp_path / "r.docx"))
    headings = [p.text for p in docx.Document(out).paragraphs if p.text.isupper() and len(p.text) < 30]
    assert headings == ["EDUCATION", "WORK EXPERIENCE", "SUMMARY"]


def test_html_matches_docx_sections():
    document = ResumeDocument(resume=_resume())
    html = HtmlResumeRenderer().render(document)
    titles = ["Summary", "Work Experience", "Skills", "Education", "Projects", "Certifications", "Interests"]
    positions = [html.index(f"<h2>{t}</h2>") for t in titles]
    assert positions == sorted(positions)
    assert "size: A4" in html and "Arial" in html
    assert "Aug 2024 – Present" in html and "2015 – 2019" in html
    assert "avery@example.com | +1 555 0100" in html


def test_a_lone_generic_skills_category_prints_no_label_p918(tmp_path):
    """The live nurse run (P9.10) printed "SKILLS / Skills: Epic, Cerner, ...":
    a single category named like the heading gets no label; real ones keep theirs."""
    resume = _resume()
    resume.skills = {"Skills": ["Epic", "Cerner"]}
    document = ResumeDocument(resume=resume)
    lines = [p.text for p in docx.Document(TemplateRenderer().render_ats_default(document, str(tmp_path / "r.docx"))).paragraphs]
    assert "Epic, Cerner" in lines and not any(t.startswith("Skills:") for t in lines)
    html = open(HtmlResumeRenderer().write_html(document, str(tmp_path / "r.html")), encoding="utf-8").read()
    assert "<p>Epic, Cerner</p>" in html

    resume.skills = {"Skills": ["Epic"], "Languages": ["Spanish"]}
    lines = [p.text for p in docx.Document(TemplateRenderer().render_ats_default(document, str(tmp_path / "r2.docx"))).paragraphs]
    assert "Skills: Epic" in lines and "Languages: Spanish" in lines
