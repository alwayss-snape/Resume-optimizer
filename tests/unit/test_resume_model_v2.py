"""Resume model v2 (P1.12): several roles at one company, project
sub-sections inside a job, and no placeholder company/title."""
import docx

from app.analysis.resume_normalizer import ResumeNormalizer
from app.domain.resume import Candidate, Experience, Resume, ResumeBullet, Role
from app.domain.resume_document import ResumeDocument
from app.ingestion.docx import RawBlock, RawDocument
from app.rendering.document_map import DocumentLocation
from app.rendering.html_renderer import HtmlResumeRenderer
from app.rendering.template_renderer import TemplateRenderer


def _raw(lines):
    """lines: (block_type, text, bold) tuples -> RawDocument."""
    blocks = []
    for i, (block_type, text, bold) in enumerate(lines):
        loc = DocumentLocation(section="x", paragraph_index=i, original_text=text)
        blocks.append(RawBlock(id=f"b{i}", block_type=block_type, text=text, location=loc, bold=bold))
    return RawDocument(filename="t.pdf", blocks=blocks)


def _experience(lines):
    raw = _raw([("name", "Sam Park", True), ("heading", "EXPERIENCE", True)] + lines)
    return ResumeNormalizer().normalize(raw)[0].resume.experience


def _resume_with_promotion():
    return Resume(
        candidate=Candidate(name="Sam Park"),
        experience=[Experience(
            id="exp_1", company="Acme Corp", title="Engineer II", location="Pune, India",
            start_date="Aug 2024", end_date="Present",
            roles=[Role(title="Engineer II", start_date="Aug 2024", end_date="Present"),
                   Role(title="Engineer I", start_date="Aug 2022", end_date="Aug 2024")],
            bullets=[
                ResumeBullet(id="b1", text="Built the ingestion service.", group="Data Platform"),
                ResumeBullet(id="b2", text="Cut costs by 20%.", group="Data Platform"),
                ResumeBullet(id="b3", text="Shipped a fraud model.", group="Fraud Detection"),
            ],
        )],
    )


def test_all_roles_and_bullet_groups():
    exp = _resume_with_promotion().experience[0]
    assert [r.title for r in exp.all_roles()] == ["Engineer II", "Engineer I"]
    assert [(g, [b.id for b in bs]) for g, bs in exp.bullet_groups()] == [
        ("Data Platform", ["b1", "b2"]), ("Fraud Detection", ["b3"]),
    ]
    single = Experience(id="e", company="X", title="Analyst", start_date="2020", end_date="2021")
    assert [(r.title, r.start_date) for r in single.all_roles()] == [("Analyst", "2020")]


def test_month_prefix_words_are_not_dates():
    """'Market' / 'Decision' / 'Junior' start with a month abbreviation; a
    project heading using them must not be read as a dated job line."""
    exp = _experience([
        ("paragraph", "Acme Corp\tPune, India", True),
        ("paragraph", "Engineer, Jan 2020 - Present", True),
        ("paragraph", "Market Basket Analysis", True),
        ("bullet", "Ranked product pairs by lift.", False),
        ("paragraph", "Decision Support Tools", True),
        ("bullet", "Built a what-if planner.", False),
    ])
    assert len(exp) == 1
    assert [g for g, _ in exp[0].bullet_groups()] == ["Market Basket Analysis", "Decision Support Tools"]


def test_title_line_before_company_line():
    exp = _experience([
        ("paragraph", "Data Analyst | March 2019 - May 2021", True),
        ("paragraph", "Globex Inc", True),
        ("bullet", "Automated weekly reports.", False),
    ])
    assert [(e.company, e.title) for e in exp] == [("Globex Inc", "Data Analyst")]
    assert exp[0].bullets[0].group is None


def test_role_after_bullets_stays_at_same_company():
    exp = _experience([
        ("paragraph", "Acme Corp", True),
        ("paragraph", "Engineer II, 2022 - Present", True),
        ("bullet", "Led the migration.", False),
        ("paragraph", "Engineer I, 2020 - 2022", True),
        ("bullet", "Wrote the test harness.", False),
    ])
    assert [(e.company, e.title) for e in exp] == [("Acme Corp", "Engineer II"), ("Acme Corp", "Engineer I")]


def test_no_placeholder_company_or_title():
    exp = _experience([("bullet", "Did a thing without any job header.", False)])
    assert (exp[0].company, exp[0].title) == ("", "")
    html = HtmlResumeRenderer().render(ResumeDocument(resume=Resume(candidate=Candidate(name="S"), experience=exp)))
    assert "Professional Experience" not in html and ">Role<" not in html


def test_docx_template_renders_roles_and_groups(tmp_path):
    out = TemplateRenderer().render_ats_default(_resume_with_promotion(), str(tmp_path / "r.docx"))
    paragraphs = [p for p in docx.Document(out).paragraphs if p.text]
    texts = [p.text for p in paragraphs]
    assert "Engineer II\tAug 2024 – Present" in texts
    assert "Engineer I\tAug 2022 – Aug 2024" in texts
    assert texts.index("Data Platform") < texts.index("Built the ingestion service.") < texts.index("Fraud Detection")
    assert texts.count("Acme Corp · Pune, India") == 1
    group_para = paragraphs[texts.index("Fraud Detection")]
    assert group_para.runs[0].bold


def test_html_renders_roles_and_groups():
    html = HtmlResumeRenderer().render(ResumeDocument(resume=_resume_with_promotion()))
    assert "<h3>Engineer II</h3>" in html
    assert "<span>Engineer I</span>" in html and "Aug 2022 – Aug 2024" in html
    assert "<p class='group'>Data Platform</p>" in html
    assert html.index("Data Platform") < html.index("Fraud Detection")


def test_date_range_only_takes_month_words():
    n = ResumeNormalizer()
    assert n._parse_title_and_dates("Senior Engineer 2019 - 2023") == ("Senior Engineer", "2019", "2023")
    assert n._parse_title_and_dates("Analyst, Sept. 2020 to Present") == ("Analyst", "Sept. 2020", "Present")
    assert n._parse_title_and_dates("Data Scientist II, August 2024 - Present") == ("Data Scientist II", "August 2024", "Present")
