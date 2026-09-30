"""P1.9: links (DOCX hyperlinks, PDF link annotations, URLs in text),
headline, DOCX walked in document order."""
import docx
import pymupdf
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.analysis.resume_normalizer import ResumeNormalizer
from app.analysis.structure_extractor import StructureExtractor
from app.domain.resume import Candidate, Resume
from app.domain.resume_document import ResumeDocument
from app.ingestion.docx import DocxParser, RawBlock, RawDocument
from app.ingestion.pdf import PdfParser
from app.rendering.document_map import DocumentLocation
from app.rendering.html_renderer import HtmlResumeRenderer
from app.rendering.template_renderer import TemplateRenderer


def _add_hyperlink(paragraph, text, url):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _normalize(raw):
    return ResumeNormalizer().normalize(raw)


def test_docx_tables_are_read_in_document_order(tmp_path):
    d = docx.Document()
    d.add_heading("Experience", level=1)
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Globex Inc"
    table.cell(0, 1).text = "Analyst, 2019 - 2021"
    d.add_paragraph("Automated the weekly sales report.", style="List Bullet")
    d.add_heading("Skills", level=1)
    d.add_paragraph("Python, SQL")
    path = tmp_path / "order.docx"
    d.save(path)

    raw = DocxParser().parse(str(path))
    texts = [b.text for b in raw.blocks]
    assert texts == ["Experience", "Globex Inc", "Analyst, 2019 - 2021", "Automated the weekly sales report.",
                     "Skills", "Python, SQL"]
    assert raw.blocks[1].section == "Experience"  # was "Skills" when tables were read last
    [exp] = _normalize(raw)[0].resume.experience
    assert (exp.company, exp.title) == ("Globex Inc", "Analyst")
    assert [b.text for b in exp.bullets] == ["Automated the weekly sales report."]


def test_docx_merged_cells_read_once(tmp_path):
    d = docx.Document()
    table = d.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Spans both columns"
    path = tmp_path / "merged.docx"
    d.save(path)
    assert [b.text for b in DocxParser().parse(str(path)).blocks] == ["Spans both columns"]


def test_docx_hyperlinks_become_candidate_links(tmp_path):
    d = docx.Document()
    d.add_paragraph("Morgan Lee")
    p = d.add_paragraph("morgan@example.com | ")
    _add_hyperlink(p, "LinkedIn", "https://www.linkedin.com/in/morgan-lee/")
    p.add_run(" | ")
    _add_hyperlink(p, "GitHub", "https://github.com/morganlee")
    _add_hyperlink(p, "Email", "mailto:morgan@example.com")
    path = tmp_path / "links.docx"
    d.save(path)

    raw = DocxParser().parse(str(path))
    assert raw.links[:2] == ["https://www.linkedin.com/in/morgan-lee/", "https://github.com/morganlee"]
    cand = _normalize(raw)[0].resume.candidate
    assert cand.links == ["https://www.linkedin.com/in/morgan-lee/", "https://github.com/morganlee"]
    assert cand.display_links() == ["linkedin.com/in/morgan-lee", "github.com/morganlee"]
    # The placeholder link labels aren't read as a headline.
    assert cand.headline is None


def test_docx_page_header_contact_details(tmp_path):
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "casey@example.com | +1 555 010 0199"
    d.add_paragraph("Casey Ortiz")
    d.add_heading("Summary", level=1)
    d.add_paragraph("Analyst.")
    path = tmp_path / "hdr.docx"
    d.save(path)
    cand = _normalize(DocxParser().parse(str(path)))[0].resume.candidate
    assert cand.email == "casey@example.com"
    assert cand.name == "Casey Ortiz"


def test_pdf_link_annotations_become_candidate_links(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 60), "Robin Hale", fontsize=22)
    page.insert_text((50, 90), "robin@example.com | LinkedIn", fontsize=11)
    page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(180, 80, 240, 95),
                      "uri": "https://linkedin.com/in/robinhale"})
    path = tmp_path / "links.pdf"
    doc.save(path)
    raw = PdfParser().parse(str(path))
    assert raw.links == ["https://linkedin.com/in/robinhale"]
    assert _normalize(raw)[0].resume.candidate.links == ["https://linkedin.com/in/robinhale"]


def _raw(lines):
    blocks = []
    for i, (block_type, text) in enumerate(lines):
        loc = DocumentLocation(section="Header", paragraph_index=i, original_text=text)
        blocks.append(RawBlock(id=f"b{i}", block_type=block_type, text=text, location=loc))
    return RawDocument(filename="t.pdf", blocks=blocks)


HEADER = [
    ("name", "Jordan Avery"),
    ("paragraph", "Senior Data Scientist | MLOps"),
    ("paragraph", "jordan@example.com | +91-9000000000 | Pune, India"),
    ("paragraph", "linkedin.com/in/jordan-avery | https://github.com/javery"),
    ("heading", "EXPERIENCE"),
    ("paragraph", "Acme Corp"),
    ("paragraph", "Data Scientist, 2020 - Present"),
    ("bullet", "Built churn models."),
]


def test_headline_and_text_urls_from_header():
    doc, evidence = _normalize(_raw(HEADER))
    cand = doc.resume.candidate
    assert cand.headline == "Senior Data Scientist | MLOps"
    assert cand.links == ["linkedin.com/in/jordan-avery", "https://github.com/javery"]
    assert cand.email == "jordan@example.com"  # the email's domain isn't taken as a link
    assert cand.location == "Pune, India"
    assert any(e.source_type == "summary" and e.text == cand.headline for e in evidence)


def test_header_lines_do_not_trigger_structure_extraction():
    raw = _raw(HEADER)
    doc, evidence = _normalize(raw)
    assert StructureExtractor(None).problems(doc.resume, evidence, raw) == []


def test_placeholder_contact_line_is_not_a_headline():
    doc, _ = _normalize(_raw([("name", "Jordan Avery"), ("paragraph", "LinkedIn | Email | Leetcode"),
                              ("paragraph", "Pune, India")]))
    assert doc.resume.candidate.headline is None


def test_renderers_show_headline_and_short_links(tmp_path):
    resume = Resume(candidate=Candidate(
        name="Jordan Avery", headline="Senior Data Scientist", email="j@example.com",
        links=["https://www.linkedin.com/in/jordan-avery/"],
    ))
    html = HtmlResumeRenderer().render(ResumeDocument(resume=resume))
    assert '<p class="headline">Senior Data Scientist</p>' in html
    assert "linkedin.com/in/jordan-avery" in html and "https://www." not in html
    out = TemplateRenderer().render_ats_default(resume, str(tmp_path / "r.docx"))
    texts = [p.text for p in docx.Document(out).paragraphs]
    assert texts[:2] == ["Jordan Avery", "Senior Data Scientist"]
    assert "linkedin.com/in/jordan-avery" in texts[2]
