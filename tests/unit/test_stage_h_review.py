"""Regressions from the independent review of Stage H (P8.1–P8.8): ordinary
US tech resumes must parse as they did before."""
from types import SimpleNamespace as B

import docx

from app.analysis.resume_normalizer import ResumeNormalizer
from app.ingestion.docx import DocxParser
from app.rendering.layout import date_range, format_date
from app.validation.coverage import content_coverage


def _parse(lines, tmp_path):
    d = docx.Document()
    for text, style in lines:
        if style == "li":
            d.add_paragraph(text, style="List Bullet")
        elif style == "hs":
            d.add_heading(text, level=2)
        else:
            run = d.add_paragraph().add_run(text)
            run.bold = style in ("h", "b")
    path = str(tmp_path / "r.docx")
    d.save(path)
    return ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume


def test_all_caps_sub_headings_stay_content(tmp_path):
    r = _parse([("John Smith", "b"), ("San Francisco, CA | (415) 555-0134 | john@example.com", "p"),
                ("EXPERIENCE", "h"), ("STRIPE", "h"), ("Senior Software Engineer, Jan 2021 - Present", "p"),
                ("Built payment pipelines.", "li"),
                ("PROJECTS", "h"), ("OPEN TRACE", "h"), ("A tracing library for Go services.", "p"),
                ("Implemented sampling.", "li"),
                ("EDUCATION", "h"), ("STANFORD UNIVERSITY", "h"), ("B.S. Computer Science, 2018", "p"),
                ("CERTIFICATIONS", "h"), ("AWS CERTIFIED SOLUTIONS ARCHITECT", "h")], tmp_path)
    assert r.other_sections == []
    assert [p.name for p in r.projects][:1] == ["OPEN TRACE"]
    assert [(e.institution, e.degree) for e in r.education] == [("STANFORD UNIVERSITY", "B.S. Computer Science")]
    assert [c["name"] for c in r.certifications] == ["AWS CERTIFIED SOLUTIONS ARCHITECT"]


def test_heading_style_project_names_stay_projects(tmp_path):
    r = _parse([("Maria Garcia", "b"), ("maria@example.com | 512-555-0199", "p"), ("Projects", "hs"),
                ("Budget Buddy", "hs"), ("Expense tracker used by 500 people.", "li"), ("Education", "hs"),
                ("University of Texas at Austin", "hs"), ("B.S. Computer Science, May 2017", "p")], tmp_path)
    assert r.other_sections == [] and r.projects and r.education


def test_a_sentence_mentioning_a_period_is_a_bullet_not_a_job(tmp_path):
    r = _parse([("Priya Patel", "b"), ("priya@example.com | 206-555-0188", "p"), ("EXPERIENCE", "h"),
                ("Data Engineer | Amazon | Jan 2020 - Present", "b"),
                ("Owned the orders data lake for the retail team, growing it from 2020 to 2023 into the main "
                 "source for finance and supply chain reporting", "p"),
                ("Led the 2021 - 2022 migration of 300 Airflow DAGs to managed workflows across four business units "
                 "and teams", "p")], tmp_path)
    assert [(e.company, len(e.bullets)) for e in r.experience] == [("Amazon", 2)]


def test_zip_plus_four_is_not_a_phone_and_framework_names_are_not_links(tmp_path):
    r = _parse([("David Lee", "b"), ("Senior .NET Engineer | ASP.NET Core | Azure", "p"),
                ("1450 Market St Apt 5, San Francisco, CA 94103-1234 | 415-555-0134 | dlee@example.com", "p"),
                ("SKILLS", "h"), ("C#, ASP.NET", "p")], tmp_path)
    c = r.candidate
    assert c.phone == "415-555-0134" and c.links == []
    assert c.headline == "Senior .NET Engineer | ASP.NET Core | Azure"


def test_headline_is_not_the_location(tmp_path):
    for headline in ("Software Engineer II", "Machine Learning Engineer | Generative AI", "Backend Engineer, Payments"):
        r = _parse([("Riley Shah", "b"), (headline, "p"), ("Palo Alto, CA | riley@example.com | 650-555-0142", "p"),
                    ("SKILLS", "h"), ("Python", "p")], tmp_path)
        assert (r.candidate.location, r.candidate.headline) == ("Palo Alto, CA", headline)


def test_us_and_eu_numeric_dates():
    assert format_date("05/06/2023") == "05/06/2023"  # May or June: printed as written
    assert format_date("08/31/2023") == "Aug 2023"
    assert date_range("01/09/2019", "31/08/2023") == "Sep 2019 – Aug 2023"


def test_merged_skill_categories_are_not_lost():
    blocks = [B(id="b1", text="Cloud: AWS"), B(id="b2", text="Testing: GoogleTest")]
    report = content_coverage(blocks, "Other: AWS, GoogleTest", ignore_words=["Cloud", "Testing"])
    assert report.lost == []
