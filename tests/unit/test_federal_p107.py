"""P10.7: the US Federal (USAJOBS-style) template. Each job's fields on
their own line, word for word, in USAJOBS order; MM/YYYY dates; header and
field lines count for keyword matching; the persona passes end to end."""
from unittest.mock import MagicMock

import docx

from app.analysis.cv_mode import federal_fields
from app.analysis.keyword_match import resume_sections
from app.domain.resume import Candidate, Experience, Resume
from app.services.tailor import TailorService

PERSONA = "data/eval/personas/federal"


def test_fields_one_per_line_in_usajobs_order_word_for_word():
    line = "Supervisor: John Doe, (202) 555-0100, may contact | Salary: $94,199 per year | 40 hours per week | Remote"
    assert federal_fields([line, "Series: 0343; Grade: GS-12"]) == [
        "40 hours per week", "Salary: $94,199 per year", "Series: 0343", "Grade: GS-12",
        "Supervisor: John Doe, (202) 555-0100, may contact", "Remote"]


def test_header_and_job_fields_are_read_for_keywords():
    resume = Resume(candidate=Candidate(name="A", details=["Highest Grade: GS-12 Step 4"]),
                    experience=[Experience(id="e", company="USDA", title="Analyst", details=["40 hours per week"])])
    texts = dict((label, text) for label, text in resume_sections(resume))
    assert texts["header"] == "Highest Grade: GS-12 Step 4" and texts["USDA (details)"] == "40 hours per week"


def test_the_federal_persona_end_to_end(tmp_path):
    service = TailorService(llm_client=None, keep_run=False)
    service.pdf_converter.convert_docx_to_pdf = MagicMock(return_value=None)
    out = service.tailor_resume(f"{PERSONA}/resume.docx", open(f"{PERSONA}/jd.txt").read(), str(tmp_path))
    lines = [p.text for p in docx.Document(out["docx"]).paragraphs if p.text.strip()]
    job = lines.index("Management and Program Analyst (GS-0343-12)\t10/2019 – Present")
    assert lines[job + 1:job + 5] == ["U.S. Department of Agriculture · Washington, DC", "40 hours per week",
                                      "Salary: $94,199 per year", "Supervisor: John Doe, (202) 555-0100, may contact"]
    assert "U.S. Citizen" in lines[2] and "Highest Grade: GS-12 Step 4" in lines[2]
    assert out["target_pages"] is None and not out["docx_warnings"]
    gs12 = next(r for r in out["keyword_match"].rows if r.keyword == "GS-12")
    assert gs12.found
    section = docx.Document(out["docx"]).sections[0]
    assert abs(section.page_height.inches - 11) < 0.01  # US Letter
