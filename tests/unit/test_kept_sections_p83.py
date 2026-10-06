"""P8.3: sections the model has no fields for are kept verbatim under their
own heading; header details are kept; licence lines aren't split."""
import docx
import pytest

from app.analysis.resume_normalizer import ResumeNormalizer
from app.domain.resume_document import ResumeDocument
from app.ingestion.docx import DocxParser
from app.rendering.html_renderer import HtmlResumeRenderer
from app.rendering.layout import section_order_for
from app.rendering.template_renderer import TemplateRenderer
from app.validation.coverage import docx_text


def _doc(lines, tmp_path):
    d = docx.Document()
    for kind, text in lines:
        if kind == "h":
            d.add_paragraph().add_run(text.upper()).bold = True
        elif kind == "name":
            d.add_paragraph().add_run(text).bold = True
        elif kind == "b":
            d.add_paragraph(text, style="List Bullet")
        else:
            d.add_paragraph(text)
    path = str(tmp_path / "r.docx")
    d.save(path)
    raw = DocxParser().parse(path)
    return raw, ResumeNormalizer().normalize(raw)[0]


@pytest.mark.parametrize("heading, kind", [
    ("Clinical Rotations", "other"), ("Work History", "experience"), ("Licenses", "certifications"),
    ("Academic Appointments", "experience"), ("Academic Projects", "projects"), ("Education & Training", "education"),
    ("Research Interests", "other"), ("Career Objective", "summary"), ("Hobbies", "interests"),
    ("Grants and Funding", "other"), ("Professional Journey", "experience"), ("Bar Admissions", "other"),
    ("Berufserfahrung", None),
])
def test_section_kinds(heading, kind):
    assert ResumeNormalizer.section_kind(heading) == kind


def test_unknown_and_foreign_sections_are_kept_and_rendered(tmp_path):
    raw, rdoc = _doc([("name", "Lucía Fernández"), ("p", "Madrid, España | +34 612 345 678 | l@example.es"),
                      ("h", "Experiencia Laboral"), ("p", "Enfermera, Hospital La Paz, 2020 – Actualidad"),
                      ("b", "Atención a pacientes críticos"), ("h", "Idiomas"), ("p", "Español (nativo)")], tmp_path)
    resume = rdoc.resume
    assert [s.heading for s in resume.other_sections] == ["Experiencia Laboral", "Idiomas"]
    assert [l.bullet for l in resume.other_sections[0].lines] == [False, True]
    rdoc.presentation.section_order = section_order_for(resume)
    out = TemplateRenderer().render_ats_default(rdoc, str(tmp_path / "out.docx"))
    text = docx_text(out)
    for phrase in ("EXPERIENCIA LABORAL", "Atención a pacientes críticos", "Español (nativo)"):
        assert phrase in text
    assert "Atención a pacientes críticos" in HtmlResumeRenderer().render(rdoc)


def test_rotations_volunteer_languages_and_licences(tmp_path):
    _, rdoc = _doc([("name", "Maria Gonzalez"), ("p", "Phoenix, AZ | (602) 555-0147 | m@example.com"),
                    ("h", "Professional Summary"), ("p", "Registered nurse."),
                    ("h", "Licenses & Certifications"),
                    ("p", "Registered Nurse (RN), Arizona State Board of Nursing, License #RN123456, exp. 06/2027"),
                    ("h", "Clinical Rotations"), ("p", "ICU - Mayo Clinic Hospital (120 hrs), Spring 2019"),
                    ("h", "Languages"), ("p", "English (native), Spanish (fluent)"),
                    ("h", "Volunteer"), ("p", "Free clinic nurse volunteer, Circle the City, 2020 - Present")], tmp_path)
    resume = rdoc.resume
    assert resume.experience == []  # rotations aren't jobs
    assert [c["name"] for c in resume.certifications] == [
        "Registered Nurse (RN), Arizona State Board of Nursing, License #RN123456, exp. 06/2027"]
    assert [s.heading for s in resume.other_sections] == ["Clinical Rotations", "Languages", "Volunteer"]


def test_header_details_kept_and_name_from_separated_header(tmp_path):
    _, rdoc = _doc([("name", "Zoë Dubois"),
                    ("p", "Date of birth: 02/11/1992 | Nationality: French | Lyon, France | +33 6 12 34 56 78 | z@example.fr"),
                    ("h", "Education"), ("p", "MSc Supply Chain, emlyon business school, 2019")], tmp_path)
    c = rdoc.resume.candidate
    assert c.details == ["Date of birth: 02/11/1992", "Nationality: French"]
    assert (c.location, c.phone) == ("Lyon, France", "+33 6 12 34 56 78")


def test_placeholder_contact_words_are_not_details(tmp_path):
    _, rdoc = _doc([("name", "Jordan Avery"), ("p", "LinkedIn | Email | Leetcode | +91-9000000000 | Pune, India"),
                    ("h", "Skills"), ("p", "Python")], tmp_path)
    assert rdoc.resume.candidate.details == []


def test_bold_label_inside_skills_is_not_a_section(tmp_path):
    d = docx.Document()
    d.add_paragraph().add_run("Avery Lee").bold = True
    d.add_paragraph().add_run("Skills").bold = True
    d.add_paragraph().add_run("Languages").bold = True
    d.add_paragraph("Python, SQL")
    path = str(tmp_path / "r.docx")
    d.save(path)
    resume = ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume
    assert resume.other_sections == [] and "Python" in [s for v in resume.skills.values() for s in v]


def test_verbatim_headings_keep_small_words_lower_case_p925():
    # Live academic run (P9.10): "GRANTS AND FUNDING" printed as "Grants And Funding".
    show = ResumeNormalizer._display_heading
    assert show("GRANTS AND FUNDING") == "Grants and Funding"
    assert show("HONORS & AWARDS:") == "Honors & Awards"
    assert show("DEAN'S LIST") == "Dean's List"
    assert show("OF NOTE") == "Of Note"  # the first word is always capitalised
    assert show("ÉDUCATION CONTINUE") == "Éducation Continue"
    assert show("Grants and Funding") == "Grants and Funding"  # mixed case is the user's own
