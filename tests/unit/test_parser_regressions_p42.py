"""Regression tests from the independent review of P4.2: inputs outside the
eval set that the first version of its parser fixes broke."""
import docx
import pymupdf
import pytest

from app.analysis.jd_analyzer import JDAnalyzer
from app.analysis.keyword_match import KeywordMatcher
from app.analysis.resume_normalizer import ResumeNormalizer
from app.domain.job import JobDescription, Requirement
from app.domain.resume import Candidate, Experience, Resume, ResumeBullet
from app.ingestion.docx import DocxParser
from app.ingestion.pdf import PdfParser


def _bold(doc, text):
    doc.add_paragraph().add_run(text).bold = True


def _parse(doc, tmp_path):
    path = str(tmp_path / "r.docx")
    doc.save(path)
    return ResumeNormalizer().normalize(DocxParser().parse(path))[0].resume


@pytest.mark.parametrize("heading", [
    "Internship Experience", "Teaching Experience", "Work Experience (Selected)", "About Me",
])
def test_experience_style_headings_still_switch_section(heading, tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "EDUCATION")
    _bold(doc, "State University")
    doc.add_paragraph("B.S. in Physics\t2016 - 2020")
    _bold(doc, heading)
    doc.add_paragraph("Data Intern | Acme Corp | Jun 2021 - Aug 2021")
    doc.add_paragraph("Built dashboards for 3 teams.", style="List Bullet")
    resume = _parse(doc, tmp_path)
    if heading == "About Me":
        assert len(resume.education) == 1  # a summary-type section; the job line isn't education
    else:
        assert [(e.company, e.title) for e in resume.experience] == [("Acme Corp", "Data Intern")]
        assert len(resume.education) == 1


@pytest.mark.parametrize("heading", ["Programming Skills", "Skills/Tools", "Soft Skills", "Licenses & Certifications",
                                     "Awards & Recognition", "Hobbies & Interests"])
def test_other_headings_still_switch_section(heading, tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "EDUCATION")
    _bold(doc, "State University")
    doc.add_paragraph("B.S. in Physics\t2016 - 2020")
    _bold(doc, heading)
    doc.add_paragraph("Python, SQL")
    resume = _parse(doc, tmp_path)
    assert len(resume.education) == 1  # the line after the heading isn't swallowed by education


def test_school_name_heading_does_not_switch_section(tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "EDUCATION")
    _bold(doc, "Riverside Institute of Technology")
    doc.add_paragraph("B.Tech in Electrical Engineering\t2014 - 2018")
    resume = _parse(doc, tmp_path)
    assert [(e.institution, e.degree) for e in resume.education] == [
        ("Riverside Institute of Technology", "B.Tech in Electrical Engineering")]
    assert not resume.skills


def test_side_by_side_skills_table_is_not_merged(tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "SKILLS")
    table = doc.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Python", "SQL, Excel"), ("Machine Learning", "Data Visualization"),
                                ("Tableau", "Power BI")]):
        table.cell(i, 0).text, table.cell(i, 1).text = a, b
    skills = {s for items in _parse(doc, tmp_path).skills.values() for s in items}
    assert {"Python", "SQL", "Excel", "Tableau", "Power BI"} <= skills


def test_labelled_skills_table_is_merged(tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "SKILLS")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Languages", "Python, SQL"
    assert _parse(doc, tmp_path).skills == {"Languages": ["Python", "SQL"]}


def test_two_column_pdf_rows_are_not_stitched(tmp_path):
    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    rows = [("Python, SQL", "Data Analyst, Globex"), ("Pune, India", "2019 - 2022"),
            ("Excel", "Built weekly revenue reports for the finance team.")]
    for i, (side, main) in enumerate(rows):
        y = 100 + i * 20
        page.insert_text((40, y), side, fontsize=10)
        page.insert_text((230, y), main, fontsize=10)
    path = str(tmp_path / "twocol.pdf")
    pdf.save(path)
    texts = [b.text for b in PdfParser().parse(path).blocks]
    assert not any("\t" in t and t.split("\t")[0] in ("Python, SQL", "Pune, India") for t in texts)


@pytest.mark.parametrize("line, found", [
    ("Led the go-to-market plan.", "Go"), ("Reported R-squared for each model.", "R"),
    ("Presented to the C-suite.", "C"), ("Joined the ml-ops guild.", "ML"), ("Chose a no-sql store.", "SQL"),
])
def test_short_keywords_do_not_match_inside_hyphenated_words(line, found):
    resume = Resume(candidate=Candidate(name="A"), experience=[
        Experience(id="e", company="X", title="Y", bullets=[ResumeBullet(id="b", text=line)])])
    job = JobDescription(keywords=[found], requirements=[Requirement(id="r", text=found)], raw_text=found)
    assert not KeywordMatcher().match(job, resume).rows[0].found


def test_term_with_a_known_suffix_still_matches():
    resume = Resume(candidate=Candidate(name="A"), experience=[
        Experience(id="e", company="X", title="Y", bullets=[ResumeBullet(id="b", text="Built HIPAA-compliant ETL.")])])
    job = JobDescription(keywords=["HIPAA"], requirements=[Requirement(id="r", text="HIPAA")], raw_text="HIPAA")
    assert KeywordMatcher().match(job, resume).rows[0].found


def test_plain_english_words_are_not_tech_keywords():
    text = "Work with the rest of the team. In spring we express our goals and excel at delivery."
    lowered = {k.lower() for k in JDAnalyzer().extract_keywords_from_text(text)}
    assert not lowered & {"rest", "spring", "express", "excel"}


@pytest.mark.parametrize("line, company, title", [
    ("Data Analyst | Lead Bank | 2020 - 2022", "Lead Bank", "Data Analyst"),
    ("Staff Engineer | Principal Financial Group | 2019 - 2023", "Principal Financial Group", "Staff Engineer"),
    ("Senior Manager - Data Platform | Globex | 2020 - 2023", "Globex", "Senior Manager"),
    ("Acme Corp | Software Engineer | 2018 - 2020", "Acme Corp", "Software Engineer"),
])
def test_title_and_company_split(line, company, title, tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "EXPERIENCE")
    doc.add_paragraph(line)
    doc.add_paragraph("Did the work.", style="List Bullet")
    exp = _parse(doc, tmp_path).experience[0]
    assert (exp.company, exp.title) == (company, title)


# -- second review ---------------------------------------------------------------

@pytest.mark.parametrize("heading", ["University Projects", "College Activities", "University Education"])
def test_headings_naming_a_school_type_still_switch(heading, tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "SKILLS")
    doc.add_paragraph("Languages: Python, SQL")
    _bold(doc, heading)
    doc.add_paragraph("Something done at university.")
    skills = {s for items in _parse(doc, tmp_path).skills.values() for s in items}
    assert skills == {"Python", "SQL"}  # the line after the heading left the skills section


@pytest.mark.parametrize("line, company, title", [
    ("Software Engineer | Google | Mountain View | 2020 - 2022", "Google", "Software Engineer"),
    ("Google | Software Engineer | Mountain View | 2020 - 2022", "Google", "Software Engineer"),
    ("Data Analyst — Payments Team — Acme Corp | 2020 - 2022", "Acme Corp", "Data Analyst"),
])
def test_location_or_team_is_not_the_company(line, company, title, tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "EXPERIENCE")
    doc.add_paragraph(line)
    doc.add_paragraph("Did the work.", style="List Bullet")
    exp = _parse(doc, tmp_path).experience[0]
    assert (exp.company, exp.title) == (company, title)


def test_lone_date_in_a_two_column_pdf_is_not_attached(tmp_path):
    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((40, 100), "Automated weekly reporting pipelines", fontsize=10)
    page.insert_text((340, 100), "2019 - 2021", fontsize=10)
    page.insert_text((40, 120), "Sidebar skill list", fontsize=10)
    page.insert_text((300, 120), "Main column text that is fairly long here", fontsize=10)
    path = str(tmp_path / "twocol2.pdf")
    pdf.save(path)
    assert not any("\t2019 - 2021" in b.text for b in PdfParser().parse(path).blocks)


def test_skills_table_header_row_is_not_merged(tmp_path):
    doc = docx.Document()
    _bold(doc, "Avery Lee")
    _bold(doc, "SKILLS")
    table = doc.add_table(rows=2, cols=2)
    for c, text in enumerate(("Languages", "Tools")):
        table.cell(0, c).paragraphs[0].add_run(text).bold = True
    table.cell(1, 0).text, table.cell(1, 1).text = "Python", "Docker"
    skills = _parse(doc, tmp_path).skills
    assert "Languages" not in skills  # never "Languages: Tools"


def test_react_and_git_need_capitals():
    lowered = {k.lower() for k in JDAnalyzer().extract_keywords_from_text("We react quickly and use git daily.")}
    assert not lowered & {"react", "git"}
