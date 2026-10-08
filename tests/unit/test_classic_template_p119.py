"""P11.9: the Classic template, an option that must read back like the standard one."""
import os

import docx
import pytest

from app.domain.resume_document import ResumePresentation, apply_style
from app.rendering.html_renderer import HtmlResumeRenderer
from app.rendering.template_renderer import TemplateRenderer, metric_runs
from app.services.tailor import TailorService
from app.validation.output import OutputQAValidator


def test_figures_are_bold_and_years_are_not():
    bold = [t for t, b in metric_runs("Scored 12M users on 6 actions, cut 10–20% in 2023; $40K, 2M+, ~12%") if b]
    assert bold == ["12M", "6", "10–20%", "$40K", "2M+", "~12%"]
    assert "".join(t for t, _ in metric_runs("Built 12M rows")) == "Built 12M rows"


def test_standard_stays_the_default_and_classic_sets_a_serif():
    p = ResumePresentation()
    assert p.style == "standard" and p.font_family == "Arial"
    apply_style(p, "classic")
    assert p.style == "classic" and p.font_family == "Cambria"
    apply_style(p, "poster")  # unknown: unchanged
    assert p.style == "classic"


@pytest.mark.parametrize("path", ["tests/fixtures/resumes/replica_layout.pdf", "tests/fixtures/resumes/sample.docx"])
def test_classic_reads_back_cleanly(tmp_path, path):
    _, doc, _ = TailorService(llm_client=None).parse_resume(path)
    apply_style(doc.presentation, "classic")
    out = str(tmp_path / "classic.docx")
    TemplateRenderer().render_ats_default(doc, out)
    assert OutputQAValidator().round_trip(out, doc.resume) == []
    paras = docx.Document(out).paragraphs
    company = next(p for p in paras if p.text.startswith(doc.resume.experience[0].company))
    assert company.runs[0].bold  # the company comes first, in bold
    assert any(r.bold and any(c.isdigit() for c in r.text) for p in paras for r in p.runs if p.style.name == "List Bullet")
    html = HtmlResumeRenderer().render(doc)
    assert 'class="classic"' in html and "serif" in html
