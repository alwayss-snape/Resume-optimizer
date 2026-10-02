"""Content coverage (P8.2): does every line of the uploaded resume reach the
output file?

The ATS round-trip (P2.5) compares the rendered file with the parsed Resume,
so anything the parser dropped was invisible to it: a German CV rendered as
just the name and email still "passed". This check compares the output with
the *source* lines instead.

A source line counts as kept when most of its words appear in the output.
Words, not exact text: the template rewrites dates ("03/2021" -> "Mar 2021"),
splits contact lines into fields and renames standard headings. Lines that
were changed on purpose are counted separately, never as lost:
- reworded: a bullet still in the resume under its own id (an accepted
  rewrite), or the summary when a tailored summary replaced it;
- trimmed: content page-fit removed to meet the page target (reported in
  its own notes).
"""
import re
from dataclasses import dataclass, field
from typing import Iterable, List, Optional, Set

# A line is kept when at least this share of its words is in the output.
KEPT_SHARE = 0.8

_STOPWORDS = {"a", "an", "the", "and", "or", "of", "in", "for", "to", "with", "at", "on", "by", "as", "from",
              "de", "la", "el", "en", "y", "und", "der", "die", "das", "von", "le", "et", "du", "des"}
_MONTHS = {"january": "jan", "february": "feb", "march": "mar", "april": "apr", "june": "jun", "july": "jul",
           "august": "aug", "september": "sep", "sept": "sep", "october": "oct", "november": "nov",
           "december": "dec"}
# Ways of writing "still in this role": the template prints "Present".
_PRESENT_RE = re.compile(r"\b(?:till date|to date|current|now|ongoing|today|heute|actualidad|aujourd'hui)\b",
                         re.IGNORECASE)


def words(text: str) -> Set[str]:
    """Comparable words of a line: lowercase, Unicode letters and digits,
    month names shortened, day / month numbers and filler words dropped."""
    text = _PRESENT_RE.sub("present", text or "")
    out = set()
    for w in re.findall(r"\w+", text.lower()):
        w = _MONTHS.get(w, w)
        if w in _STOPWORDS or w == "_":
            continue
        if w.isdigit() and len(w) <= 2:
            continue  # "03/2021" is printed "Mar 2021"
        if len(w) < 2 and not w.isdigit():
            continue
        out.add(w)
    return out


@dataclass
class CoverageReport:
    counted: int = 0
    kept: int = 0
    reworded: int = 0
    trimmed: int = 0
    lost: List[str] = field(default_factory=list)

    @property
    def pct(self) -> Optional[float]:
        if not self.counted:
            return None
        return round(100.0 * (self.counted - len(self.lost)) / self.counted, 1)

    def as_dict(self) -> dict:
        return {"pct": self.pct, "counted": self.counted, "kept": self.kept, "reworded": self.reworded,
                "trimmed": self.trimmed, "lost": list(self.lost)}


def content_coverage(blocks: Iterable, output_text: str, *, heading_texts: Iterable[str] = (),
                     reworded_blocks: Iterable[str] = (), trimmed_blocks: Iterable[str] = (),
                     removed_text: str = "", ignore_words: Iterable[str] = ()) -> CoverageReport:
    """`blocks`: the source RawBlocks. `heading_texts`: standard section
    headings, which the template renames, so they aren't counted.
    `reworded_blocks` / `trimmed_blocks`: block ids changed or removed on
    purpose. `removed_text`: text page-fit removed (a dropped Interests line
    has no block id of its own to point at). `ignore_words`: labels the
    template may merge away (skill categories beyond four become "Other")."""
    have = words(output_text)
    ignored = set().union(*(words(w) for w in ignore_words)) if ignore_words else set()
    removed = words(removed_text)
    headings = {re.sub(r"\s+", " ", h).strip().lower() for h in heading_texts}
    reworded, trimmed = set(reworded_blocks), set(trimmed_blocks)
    report = CoverageReport()
    for block in blocks:
        text = (block.text or "").strip()
        if not text or re.sub(r"\s+", " ", text).lower() in headings:
            continue
        line_words = words(text) - ignored if ignored else words(text)
        if not line_words:
            continue
        report.counted += 1
        if block.id in reworded:
            report.reworded += 1
            continue
        if block.id in trimmed:
            report.trimmed += 1
            continue
        share = len(line_words & have) / len(line_words)
        if share >= KEPT_SHARE:
            report.kept += 1
        elif removed and len(line_words & (have | removed)) / len(line_words) >= KEPT_SHARE:
            report.trimmed += 1
        else:
            report.lost.append(text)
    return report


def docx_text(path: Optional[str]) -> str:
    """All text of a DOCX: body paragraphs and table cells."""
    import os
    if not path or not os.path.exists(path):
        return ""
    import docx
    d = docx.Document(path)
    cells = [c.text for t in d.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in d.paragraphs] + cells)
