"""Generate fictional persona resumes (DOCX + some PDFs) and JDs for user testing.
Spec lines: ("name"|"h"|"hs"|"p"|"b"|"bold"|"pb", text)
  h  = bold ALL-CAPS heading paragraph (most common in real resumes)
  hs = Word "Heading 1" style
  b  = "List Bullet" style bullet
  pb = paragraph starting with a glyph bullet
  bold = bold line (company/title)
"""
import os, sys
import docx
from docx.shared import Pt, Inches

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R = os.path.join(BASE, "resumes"); J = os.path.join(BASE, "jds")

def write_docx(path, spec, header=None, footer=None):
    d = docx.Document()
    if header:
        d.sections[0].header.paragraphs[0].text = header
    if footer:
        d.sections[0].footer.paragraphs[0].text = footer
    for kind, text in spec:
        if kind == "name":
            p = d.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(18)
        elif kind == "h":
            p = d.add_paragraph(); r = p.add_run(text.upper()); r.bold = True
        elif kind == "hs":
            d.add_heading(text, level=1)
        elif kind == "b":
            d.add_paragraph(text, style="List Bullet")
        elif kind == "pb":
            d.add_paragraph("• " + text)
        elif kind == "bold":
            p = d.add_paragraph(); r = p.add_run(text); r.bold = True
        else:
            d.add_paragraph(text)
    d.save(path)

def write_pdf(path, spec):
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    y = 50
    def line(text, size=10, bold=False, x=50):
        nonlocal page, y
        if y > 800:
            page = doc.new_page(width=595, height=842); y = 50
        page.insert_text((x, y), text, fontsize=size, fontname="hebo" if bold else "helv")
        y += size + 4
    for kind, text in spec:
        if kind == "name": line(text, 18, True)
        elif kind in ("h", "hs"): y += 4; line(text.upper(), 12, True)
        elif kind == "bold": line(text, 10, True)
        elif kind in ("b", "pb"):
            # naive wrap at ~95 chars
            words, cur, first = text.split(), "", True
            for w in words:
                if len(cur) + len(w) > 95:
                    line(("• " if first else "  ") + cur, x=60 if first else 68); cur, first = w, False
                else:
                    cur = (cur + " " + w).strip()
            line(("• " if first else "  ") + cur, x=60 if first else 68)
        else: line(text)
    doc.save(path)

P = {}
JD = {}

# 1. Registered nurse
P["nurse"] = [
 ("name","Maria Gonzalez, RN, BSN"),
 ("p","Phoenix, AZ | (602) 555-0147 | maria.gonzalez.rn@example.com"),
 ("h","Professional Summary"),
 ("p","Compassionate Registered Nurse with 6 years of acute care experience in med-surg and step-down units. Skilled in patient assessment, care planning and family education."),
 ("h","Licenses & Certifications"),
 ("p","Registered Nurse (RN), Arizona State Board of Nursing, License #RN123456, exp. 06/2027 (Compact/NLC)"),
 ("p","BLS - American Heart Association, exp. 03/2026"),
 ("p","ACLS - American Heart Association, exp. 03/2026"),
 ("p","PALS, exp. 01/2026"),
 ("h","Clinical Experience"),
 ("bold","Registered Nurse - Step-Down Unit\tBanner University Medical Center, Phoenix, AZ"),
 ("p","03/2021 - Present"),
 ("b","Provide direct patient care for 4-5 telemetry patients per 12-hour night shift"),
 ("b","Titrate cardiac drips (diltiazem, heparin, amiodarone) per protocol and monitor rhythms"),
 ("b","Precept new graduate nurses; trained 8 orientees on Epic charting and unit workflows"),
 ("b","Charge nurse 2x per month, coordinating assignments for 24-bed unit"),
 ("bold","Registered Nurse - Medical-Surgical\tSt. Joseph's Hospital, Phoenix, AZ"),
 ("p","06/2019 - 02/2021"),
 ("b","Cared for 5-6 post-operative patients per shift, including wound care, pain management and discharge teaching"),
 ("b","Member of unit falls-prevention committee; unit falls decreased 30% in 2020"),
 ("h","Clinical Rotations"),
 ("p","ICU - Mayo Clinic Hospital (120 hrs), Spring 2019"),
 ("p","Pediatrics - Phoenix Children's (90 hrs), Fall 2018"),
 ("h","Education"),
 ("p","Bachelor of Science in Nursing (BSN), Arizona State University, 2019"),
 ("h","Skills"),
 ("p","Epic, Cerner, telemetry monitoring, IV insertion, wound care, Spanish (fluent)"),
 ("h","Languages"),
 ("p","English (native), Spanish (fluent)"),
 ("h","Volunteer"),
 ("p","Free clinic nurse volunteer, Circle the City, 2020 - Present"),
]
JD["nurse"] = """Registered Nurse - ICU (Nights)
HonorHealth Scottsdale Osborn Medical Center

We are seeking a compassionate ICU Registered Nurse to join our critical care team.

Responsibilities:
- Provide direct patient care to critically ill patients using the nursing process
- Monitor hemodynamics, titrate vasoactive drips, manage ventilated patients
- Collaborate with interdisciplinary team; educate patients and families
- Document care accurately in Epic

Requirements:
- Current RN license in Arizona or compact (NLC) state
- BLS and ACLS certification required; CCRN preferred
- Minimum 2 years of acute care experience; 1 year ICU preferred
- BSN required
- Ability to work 12-hour night shifts, weekends and holidays
- Ability to lift 50 lbs and stand for extended periods
- Bilingual English/Spanish a plus
- Knowledge of HIPAA regulations

HonorHealth is an Equal Opportunity Employer. Benefits include medical, dental, 403(b), tuition reimbursement.
"""

# 2. Teacher -> instructional designer
P["teacher"] = [
 ("name","Priya Raman"),
 ("p","Austin, TX · priya.raman@example.com · 512-555-0199 · portfolio: priyaraman.design"),
 ("h","Professional Journey"),
 ("bold","8th Grade Science Teacher | Round Rock ISD | Aug 2016 - Jun 2024"),
 ("b","Designed and taught inquiry-based science curriculum for 150 students per year"),
 ("b","Built unit plans and assessments in Google Classroom and Canvas aligned to TEKS standards"),
 ("b","Led professional development sessions for 25 teachers on differentiated instruction"),
 ("b","Raised state STAAR science pass rate from 71% to 84% over three years"),
 ("bold","Curriculum Writer (contract) | Texas Education Agency | Summer 2021"),
 ("b","Wrote 12 lesson modules for statewide virtual learning program"),
 ("h","Instructional Design Training"),
 ("p","ATD Instructional Design Certificate, 2024"),
 ("p","Self-taught: Articulate Rise 360, Camtasia, Canva"),
 ("h","Education"),
 ("p","M.Ed. Curriculum and Instruction, Texas State University, 2019"),
 ("p","B.S. Biology, University of Texas at Austin, 2016"),
]
JD["teacher"] = """Instructional Designer
Acme Learning Co. is looking for an Instructional Designer to build engaging e-learning for our corporate clients.

What you'll do
• Design and develop e-learning modules using Articulate Storyline and Rise 360
• Apply ADDIE and SAM models and adult learning theory
• Partner with SMEs to write storyboards, scripts and assessments
• Publish SCORM packages to the LMS and analyze learner data

What you'll bring
• 3+ years of instructional design experience
• Portfolio of e-learning work required
• Experience with Articulate Storyline, Camtasia, and an LMS (Cornerstone or Docebo)
• Strong stakeholder management and project management skills
• Knowledge of WCAG accessibility standards
"""

# 3. Electrician (short, no metrics)
P["electrician"] = [
 ("name","DEREK O'BRIEN"),
 ("p","Cleveland OH   216-555-0110   derekobrien77@example.com"),
 ("h","Work History"),
 ("p","Journeyman Electrician, Lakeshore Electric LLC, 2017 - Current"),
 ("pb","Install and repair wiring, conduit, panels in commercial buildings"),
 ("pb","Read blueprints and follow NEC code"),
 ("pb","Troubleshoot motors and controls"),
 ("p","Apprentice Electrician, IBEW Local 38, 2013 - 2017"),
 ("pb","Assisted journeymen on industrial and residential jobs"),
 ("h","Licenses"),
 ("p","Ohio Journeyman Electrician License"),
 ("p","OSHA 30"),
 ("p","Valid Ohio driver's license, clean record"),
 ("h","Training"),
 ("p","IBEW/NECA 5-year apprenticeship program, 2017"),
]
JD["electrician"] = """Industrial Maintenance Electrician - 2nd shift
Pay: $36-$42/hr
We need an experienced industrial electrician for our manufacturing plant.
Requirements:
* Journeyman electrician license
* 5+ years industrial electrical experience
* PLC troubleshooting (Allen-Bradley ControlLogix) and VFDs
* Knowledge of NEC and NFPA 70E, lockout/tagout (LOTO)
* OSHA 10 or 30
* Must be able to lift 50 lbs, climb ladders and work at heights
* Available for overtime and weekends
* High school diploma or GED
"""

# 4. Warehouse/logistics
P["warehouse"] = [
 ("name","Tyrone Jackson"),
 ("p","Memphis, TN | 901-555-0123 | tjackson@example.com"),
 ("h","Experience"),
 ("bold","Warehouse Associate — FedEx Ground Hub, Memphis TN — 2019 to present"),
 ("b","Load and unload trailers, scan packages with RF scanner"),
 ("b","Operate forklift, reach truck and pallet jack"),
 ("b","Trained new hires on safety procedures"),
 ("bold","Delivery Driver (part-time, concurrent) — DoorDash — 2020 to present"),
 ("b","Complete 40+ deliveries a week with 4.9 customer rating"),
 ("h","Certifications"),
 ("p","Forklift Operator Certification (OSHA), CDL Class A (2022), HAZMAT endorsement"),
 ("h","Education"),
 ("p","High School Diploma, Whitehaven High School, 2017"),
]
JD["warehouse"] = """Logistics Coordinator
Responsibilities: coordinate inbound and outbound shipments; manage carrier relationships; track inventory in WMS (Manhattan); prepare BOLs; ensure DOT compliance.
Requirements: 2+ years warehouse or logistics experience; CDL a plus; proficiency in Microsoft Excel; strong communication skills; forklift certification preferred; ability to work in a fast-paced environment.
"""

# 5. Retail / hospitality
P["retail"] = [
 ("name","Aaliyah Brooks"),
 ("p","Atlanta, GA · (404) 555-0188 · aaliyah.brooks@example.com"),
 ("h","Summary"),
 ("p","Customer-focused shift supervisor with 5 years in retail and food service."),
 ("h","Experience"),
 ("bold","Shift Supervisor, Starbucks, Atlanta GA, 2021 - Present"),
 ("b","Lead a team of 6-8 baristas per shift; open and close store, handle cash reconciliation"),
 ("b","Coach partners on customer service; store Customer Connection score in top 10% of district"),
 ("b","Manage inventory orders and weekly scheduling"),
 ("bold","Sales Associate, Target, Atlanta GA, 2019 - 2021"),
 ("b","Assisted guests, processed returns, stocked shelves, set up planograms"),
 ("h","Education"),
 ("p","Associate of Arts, Georgia State University Perimeter College, 2019"),
 ("h","Skills"),
 ("p","POS systems, cash handling, scheduling, conflict resolution, team leadership"),
]
JD["retail"] = """Assistant Store Manager - Retail
Join our team! We're hiring an Assistant Store Manager.
- 2+ years retail management or supervisory experience
- Drive sales and achieve store KPIs; analyze P&L
- Recruit, train and develop team members
- Merchandising, visual standards and planograms
- Inventory control, loss prevention and shrink reduction
- Excellent customer service and communication skills
- Flexible availability including nights, weekends, holidays
- Proficiency with POS and Microsoft Office
"""

# 6. Sales AE
P["sales"] = [
 ("name","Jordan Kim"),
 ("p","Chicago, IL | jordan.kim@example.com | 312-555-0166 | linkedin.com/in/jordankim"),
 ("h","Summary"),
 ("p","Quota-crushing Account Executive with 7 years of B2B SaaS sales."),
 ("h","Professional Experience"),
 ("bold","Senior Account Executive, CloudMetrics Inc., Chicago, IL\tJan 2021 – Present"),
 ("b","Achieved 132% of $1.2M annual quota in FY2023 and 118% in FY2022"),
 ("b","Closed largest deal in company history ($480K ACV) with a Fortune 500 retailer"),
 ("b","Ran full sales cycle from prospecting to close using Salesforce and Outreach"),
 ("bold","Account Executive, PayStream, Chicago, IL\tMar 2018 – Dec 2020"),
 ("b","Ranked #2 of 14 AEs in 2019; President's Club 2019"),
 ("b","Built pipeline of $3M through outbound cold calling and LinkedIn Sales Navigator"),
 ("bold","Sales Development Representative, PayStream\tJun 2016 – Feb 2018"),
 ("b","Booked 25+ qualified meetings per month"),
 ("h","Education"),
 ("p","B.A. Communications, University of Illinois, 2016"),
]
JD["sales"] = """Enterprise Account Executive
About the role: You will own a $1.5M+ quota selling our platform to enterprise accounts.
Requirements
- 5+ years of closing experience in B2B SaaS, with a track record of exceeding quota
- Experience with MEDDIC or MEDDPICC sales methodology
- Proficiency in Salesforce CRM and Gong
- Strong negotiation, forecasting and pipeline management skills
- Experience selling to C-level executives
- Bachelor's degree preferred
Benefits: uncapped commission, equity, unlimited PTO.
"""

# 7. Accountant CPA
P["accountant"] = [
 ("name","Rebecca Liu, CPA"),
 ("p","New York, NY | rebecca.liu@example.com | 646-555-0177"),
 ("h","Profile"),
 ("p","CPA with 5 years of public accounting and industry experience in month-end close, financial reporting and SOX compliance."),
 ("h","Experience"),
 ("bold","Senior Accountant, Brightline Media, New York, NY\t07/2022 – Present"),
 ("b","Own month-end close for 3 entities; reduced close from 8 to 5 business days"),
 ("b","Prepare journal entries, account reconciliations and accruals under US GAAP"),
 ("b","Built Excel models with pivot tables and VLOOKUP/INDEX-MATCH for variance analysis"),
 ("bold","Audit Associate, KPMG LLP, New York, NY\t09/2019 – 06/2022"),
 ("b","Performed audit procedures for clients in media and consumer goods; tested SOX 404 controls"),
 ("h","Certifications"),
 ("p","Certified Public Accountant (CPA), New York, 2021"),
 ("h","Education"),
 ("p","Master of Accountancy (MAcc), Baruch College, 2019"),
 ("h","Technical Skills"),
 ("p","NetSuite, SAP, Excel (advanced), BlackLine, Power BI"),
]
JD["accountant"] = """Senior Accountant
Responsibilities: Manage the month-end and year-end close process; prepare financial statements in accordance with Generally Accepted Accounting Principles; perform balance sheet reconciliations; support external audit; maintain internal controls (SOX).
Qualifications: Bachelor's degree in Accounting; CPA required; 4+ years of accounting experience; Big 4 experience preferred; advanced Microsoft Excel; experience with NetSuite or Oracle ERP; knowledge of ASC 606 revenue recognition.
"""

# 8. Lawyer / paralegal
P["lawyer"] = [
 ("name","Samuel Adeyemi, Esq."),
 ("p","Washington, DC | samuel.adeyemi@example.com | 202-555-0135"),
 ("h","Bar Admissions"),
 ("p","District of Columbia Bar, 2018; New York State Bar, 2017"),
 ("h","Legal Experience"),
 ("bold","Associate, Hale & Morgan LLP, Washington, DC\tSep 2018 – Present"),
 ("b","Draft motions, briefs and discovery requests in federal commercial litigation matters"),
 ("b","Took and defended 20+ depositions; second-chaired two jury trials"),
 ("b","Manage document review teams using Relativity"),
 ("bold","Judicial Law Clerk, Hon. Ellen Park, U.S. District Court for the District of Maryland\tAug 2017 – Aug 2018"),
 ("b","Drafted bench memoranda and opinions on civil and criminal motions"),
 ("h","Publications"),
 ("p","\"Arbitration Clauses After Epic Systems,\" 98 Geo. L.J. Online 45 (2019)"),
 ("p","\"Discovery of Ephemeral Messaging,\" ABA Litigation Journal (2021)"),
 ("h","Education"),
 ("p","J.D., cum laude, Georgetown University Law Center, 2017"),
 ("p","B.A. Political Science, Howard University, 2014"),
 ("h","References"),
 ("p","Available upon request"),
]
JD["lawyer"] = """Litigation Associate (4th-6th year)
Our DC office seeks a commercial litigation associate.
Qualifications:
- J.D. from an accredited law school; active DC Bar membership required
- 4-6 years of complex commercial litigation experience at a law firm
- Federal court experience, including taking depositions and motion practice
- Excellent legal research and writing skills (Westlaw, Lexis)
- eDiscovery experience (Relativity) preferred
- Judicial clerkship preferred
"""

# 9. Academic CV (long)
acad = [("name","Dr. Hannah Okafor"),
 ("p","Department of Biology, Midwestern State University | h.okafor@example.edu | ORCID 0000-0002-1234-5678"),
 ("h","Research Interests"),
 ("p","Microbial ecology, soil metagenomics, climate-driven carbon cycling."),
 ("h","Academic Appointments"),
 ("bold","Assistant Professor of Biology, Midwestern State University\t2019 – Present"),
 ("b","Lead a lab of 4 PhD students and 6 undergraduates studying soil microbial communities"),
 ("b","Teach BIO 210 Microbiology (120 students) and graduate seminar in metagenomics"),
 ("bold","Postdoctoral Research Fellow, University of Wisconsin–Madison\t2016 – 2019"),
 ("b","Developed metagenomic pipelines in R and Python for 2,000 soil samples"),
 ("h","Education"),
 ("p","Ph.D. Microbiology, Cornell University, 2016"),
 ("p","B.Sc. Biochemistry, University of Lagos, 2010"),
 ("h","Grants and Funding"),
 ("p","NSF CAREER Award, PI, $850,000, 2022–2027"),
 ("p","USDA NIFA Foundational Program, Co-PI, $500,000, 2020–2023"),
 ("h","Publications"),
]
for i in range(1, 46):
    acad.append(("p", f"{i}. Okafor H., Smith J., Lee K. ({2010 + i % 14}). Soil microbial response to warming experiment {i}. Journal of Microbial Ecology {40+i}({i%12+1}): {100+i}-{115+i}. doi:10.1000/jme.{2000+i}"))
acad += [("h","Teaching"), ("p","BIO 210 Microbiology, 2019–present; BIO 590 Metagenomics Seminar, 2020–present"),
 ("h","Invited Talks"), ("p","ASM Microbe 2023, Houston TX; ISME 2022, Lausanne, Switzerland"),
 ("h","Service"), ("p","Reviewer for ISME Journal, mSystems; Graduate admissions committee 2021–2023"),
 ("h","Professional Memberships"), ("p","American Society for Microbiology; Ecological Society of America")]
P["academic"] = acad
JD["academic"] = """Senior Scientist, Soil Microbiome (R&D)
Pivot AgBio is hiring a Senior Scientist to lead soil microbiome discovery.
Requirements:
- PhD in Microbiology, Ecology, or related field with 5+ years of postdoctoral or industry experience
- Track record of peer-reviewed publications in microbial ecology
- Expertise in metagenomics, 16S rRNA amplicon sequencing and bioinformatics (Python, R, QIIME2)
- Experience managing research teams and securing grant funding
- Strong scientific communication skills
Preferred: experience with greenhouse or field trials; industry R&D experience
"""

# 10. Designer (two-column table layout, minimal text) - built separately
JD["designer"] = """Senior Product Designer
We are looking for a Senior Product Designer with 5+ years of experience designing web and mobile products.
- Expert in Figma, prototyping and design systems
- Experience conducting user research and usability testing
- Strong portfolio demonstrating end-to-end product design
- Collaborate with product managers and engineers
- Familiarity with accessibility (WCAG) and HTML/CSS a plus
"""

# 11. Executive 20+ years
ex = [("name","Robert J. Whitfield"),
 ("p","Dallas, TX | rwhitfield@example.com | 214-555-0101 | linkedin.com/in/rjwhitfield"),
 ("h","Executive Summary"),
 ("p","Operations executive with 24 years leading supply chain and manufacturing organizations of up to 3,000 employees and $1.4B P&L."),
 ("h","Career History")]
jobs = [("Chief Operating Officer","Atlas Industrial Group","2018","Present"),
        ("SVP, Global Supply Chain","Meridian Consumer Products","2013","2018"),
        ("VP Manufacturing","Meridian Consumer Products","2009","2013"),
        ("Plant Director","Kessler Foods","2005","2009"),
        ("Operations Manager","Kessler Foods","2001","2005"),
        ("Production Supervisor","General Mills","1998","2001"),
        ("Industrial Engineer","General Mills","1996","1998")]
for t,c,s,e in jobs:
    ex.append(("bold", f"{t} | {c} | {s} – {e}"))
    for k in range(4):
        ex.append(("b", f"Led {['lean transformation','S&OP redesign','ERP (SAP) rollout','capex program'][k]} across {k+3} sites, delivering ${(k+1)*12}M in annualized savings and improving OTIF by {k+4} points"))
ex += [("h","Board & Advisory"),("p","Board Member, Dallas Regional Chamber, 2019 – Present"),
       ("h","Education"),("p","MBA, Kellogg School of Management, Northwestern University, 2005"),("p","B.S. Industrial Engineering, Purdue University, 1996")]
P["executive"] = ex
JD["executive"] = """Vice President of Operations
We seek a VP of Operations to lead 6 manufacturing plants. Requirements: 15+ years of operations leadership in CPG manufacturing; P&L ownership; Lean Six Sigma (Black Belt preferred); S&OP; SAP; experience leading M&A integration; MBA preferred; strong executive presence and board communication.
"""

# 12. New grad with only education + projects, no Experience
P["newgrad"] = [
 ("name","Emily Chen"),
 ("p","emily.chen@example.edu | (617) 555-0142 | github.com/emilychen | Boston, MA"),
 ("h","Education"),
 ("p","B.S. Statistics, Boston University, Expected May 2026, GPA 3.7/4.0"),
 ("p","Relevant Coursework: Regression Analysis, Data Mining, Database Systems, Probability"),
 ("h","Academic Projects"),
 ("bold","MBTA Ridership Dashboard"),
 ("b","Built a Tableau dashboard of 5 years of MBTA ridership data using SQL and Python (pandas)"),
 ("bold","Predicting Housing Prices"),
 ("b","Trained linear regression and random forest models in scikit-learn; R² of 0.82"),
 ("h","Leadership & Activities"),
 ("p","Treasurer, BU Statistics Club (2024–2025); Teaching Assistant, MA 113 (Fall 2025)"),
 ("h","Skills"),
 ("p","Python, R, SQL, Excel, Tableau"),
]
JD["newgrad"] = """Junior Data Analyst (Entry level)
Requirements: Bachelor's degree in Statistics, Mathematics, Economics or related; proficiency in SQL and Excel; experience with Tableau or Power BI; Python or R; strong communication skills; internship experience a plus; attention to detail.
"""

# 13. Career break / freelance / concurrent jobs
P["gap"] = [
 ("name","Nicole Fischer"),
 ("p","Denver, CO | nicole.fischer@example.com | 303-555-0155"),
 ("h","Summary"),
 ("p","Marketing professional returning after a two-year career break for family caregiving."),
 ("h","Experience"),
 ("bold","Freelance Content Marketer | Self-employed | 2023 – Present"),
 ("b","Write SEO blog content and email campaigns for 5 small-business clients"),
 ("bold","Career Break — Family Caregiver | 2021 – 2023"),
 ("b","Full-time caregiver for a parent; managed medical scheduling and household finances"),
 ("bold","Marketing Manager | Summit Outdoor Co. | Mar 2016 – Feb 2021"),
 ("b","Managed $400K annual paid media budget across Google Ads and Meta"),
 ("b","Grew email list from 20K to 85K subscribers using HubSpot"),
 ("bold","Adjunct Instructor (concurrent) | Community College of Denver | 2018 – 2020"),
 ("b","Taught Intro to Marketing to 30 students per semester"),
 ("h","Education"),
 ("p","B.A. Marketing, Colorado State University, 2015"),
]
JD["gap"] = """Digital Marketing Manager
Responsibilities: own paid search and paid social campaigns; manage marketing automation in HubSpot; SEO and content strategy; report on ROI and CAC; manage agencies.
Requirements: 5+ years digital marketing experience; Google Ads and Meta Ads certification preferred; HubSpot; Google Analytics 4; strong analytical skills; budget management experience.
"""

# 14. Veteran
P["veteran"] = [
 ("name","SGT Marcus Reyes (Ret.)"),
 ("p","San Antonio, TX | marcus.reyes@example.com | 210-555-0190"),
 ("p","Security Clearance: Active Secret"),
 ("h","Military Experience"),
 ("bold","Logistics NCOIC (92Y Unit Supply Specialist), U.S. Army, Fort Hood, TX\t2014 – 2024"),
 ("b","Accountable for $12M in property book equipment with zero loss over 4 annual inventories"),
 ("b","Supervised 9 Soldiers; managed GCSS-Army transactions and hand receipts"),
 ("b","Coordinated movement of 200+ vehicles for NTC rotation"),
 ("h","Awards"),
 ("p","Army Commendation Medal (2), Army Achievement Medal (3)"),
 ("h","Education & Training"),
 ("p","Advanced Leader Course (ALC), 2019; B.S. Business Administration (in progress), UMGC"),
]
JD["veteran"] = """Supply Chain Manager - Defense Contractor
Requirements: Active Secret clearance required; 5+ years supply chain or logistics experience; inventory management and property accountability; ERP systems (SAP); team leadership; PMP or APICS CSCP preferred; Bachelor's degree.
"""

# 15. European CV (German) with DD/MM/YYYY and personal details, accents
P["eu_cv"] = [
 ("name","José Müller-Øberg"),
 ("p","Lebenslauf / Curriculum Vitae"),
 ("p","Geburtsdatum: 14.03.1990 | Staatsangehörigkeit: Deutsch | Familienstand: verheiratet"),
 ("p","Hauptstraße 12, 10115 Berlin | +49 30 12345678 | jose.mueller@example.de"),
 ("h","Berufserfahrung"),
 ("bold","Projektmanager, Siemens AG, München\t01/2019 – heute"),
 ("b","Leitung von IT-Projekten mit Budgets bis 2 Mio. €"),
 ("b","Einführung von SAP S/4HANA in 3 Werken"),
 ("bold","Berater, Accenture GmbH, Berlin\t09/2015 – 12/2018"),
 ("b","Beratung von Kunden im Bereich Prozessoptimierung"),
 ("h","Ausbildung"),
 ("p","M.Sc. Wirtschaftsinformatik, Technische Universität München, 2015"),
 ("h","Sprachen"),
 ("p","Deutsch (Muttersprache), Englisch (C1), Spanisch (B2)"),
 ("h","Kenntnisse"),
 ("p","SAP S/4HANA, Jira, Scrum, MS Project"),
]
JD["eu_cv"] = """Projektmanager (m/w/d) IT
Ihre Aufgaben:
- Leitung von IT-Projekten im SAP-Umfeld
- Budget- und Ressourcenplanung
Ihr Profil:
- Abgeschlossenes Studium der Wirtschaftsinformatik
- Mehrjährige Erfahrung im Projektmanagement
- Kenntnisse in SAP S/4HANA und agilen Methoden (Scrum)
- PMP- oder PRINCE2-Zertifizierung wünschenswert
- Sehr gute Deutsch- und Englischkenntnisse
"""

# 16. European CV in English, DD/MM/YYYY dates
P["eu_en"] = [
 ("name","Zoë Dubois"),
 ("p","Curriculum Vitae"),
 ("p","Date of birth: 02/11/1992 | Nationality: French | Lyon, France | +33 6 12 34 56 78 | zoe.dubois@example.fr"),
 ("h","Work Experience"),
 ("bold","Supply Chain Analyst, Groupe SEB, Lyon\t01/09/2019 – 31/08/2023"),
 ("b","Analysed demand forecasts with SAP APO and Excel for 1,200 SKUs"),
 ("bold","Logistics Intern, Danone, Paris\t03/2018 – 08/2018"),
 ("b","Optimised warehouse slotting, cutting picking time 12%"),
 ("h","Education"),
 ("p","MSc Supply Chain Management, emlyon business school, 2019"),
 ("h","Languages"),
 ("p","French (native), English (C1), German (B1)"),
 ("h","Hobbies"),
 ("p","Trail running, photography"),
]
JD["eu_en"] = """Supply Chain Analyst (EMEA)
Requirements: 3+ years in supply chain planning; SAP APO or IBP; advanced Excel; demand forecasting; fluent English and French; German is a plus; Power BI.
"""

# 17. Spanish resume + Spanish JD
P["spanish"] = [
 ("name","Lucía Fernández García"),
 ("p","Madrid, España | +34 612 345 678 | lucia.fernandez@example.es"),
 ("h","Perfil Profesional"),
 ("p","Enfermera con 5 años de experiencia en atención primaria y urgencias."),
 ("h","Experiencia Laboral"),
 ("bold","Enfermera de Urgencias, Hospital Universitario La Paz, Madrid\t2020 – Actualidad"),
 ("b","Atención a pacientes críticos en el servicio de urgencias"),
 ("b","Administración de medicación y cuidados de heridas"),
 ("h","Formación Académica"),
 ("p","Grado en Enfermería, Universidad Complutense de Madrid, 2019"),
 ("h","Idiomas"),
 ("p","Español (nativo), Inglés (B2)"),
]
JD["spanish"] = """Enfermera/o de Cuidados Intensivos
Requisitos:
- Grado en Enfermería y colegiación vigente
- Experiencia mínima de 2 años en urgencias o UCI
- Conocimientos de soporte vital avanzado (SVA)
- Manejo de ventilación mecánica
- Disponibilidad para turnos rotativos
"""

# 18. CJK name, Indian-format resume
P["india"] = [
 ("name","王小明 (Wang Xiaoming)"),
 ("p","Bengaluru, India | +91 98765 43210 | wang.xm@example.com"),
 ("h","Career Objective"),
 ("p","To obtain a challenging position in a reputed organisation."),
 ("h","Work Experience"),
 ("bold","Senior Software Engineer, Infosys Ltd, Bengaluru\tJuly 2018 – Till Date"),
 ("b","Developed REST APIs in Java Spring Boot for a banking client"),
 ("bold","Software Engineer, TCS, Chennai\tJune 2015 – June 2018"),
 ("b","Maintained Oracle PL/SQL procedures"),
 ("h","Educational Qualifications"),
 ("p","B.E. Computer Science, Anna University, 2015 — 78%"),
 ("h","Personal Details"),
 ("p","Father's Name: Wang Wei; Date of Birth: 05-08-1993; Marital Status: Single; Passport: Available"),
 ("h","Declaration"),
 ("p","I hereby declare that the above information is true to the best of my knowledge."),
]
JD["india"] = """Senior Java Developer
Requirements: 6+ years Java, Spring Boot, microservices, REST APIs, Kafka, AWS, Oracle SQL; banking domain experience preferred; notice period up to 30 days.
"""

# 19. Federal resume
P["federal"] = [
 ("name","Angela Martin"),
 ("p","1234 Elm St, Arlington, VA 22201 | angela.martin@example.com | 703-555-0129"),
 ("p","U.S. Citizen | Veterans' Preference: None | Highest Grade: GS-12 Step 4"),
 ("h","Work Experience"),
 ("bold","Management and Program Analyst (GS-0343-12)"),
 ("p","U.S. Department of Agriculture, Washington, DC"),
 ("p","10/2019 – Present | 40 hours per week | Salary: $94,199 per year | Supervisor: John Doe, (202) 555-0100, may contact"),
 ("b","Conduct program evaluations and prepare briefing materials for senior leadership"),
 ("b","Manage $3M budget execution and track obligations in FMMI"),
 ("h","Education"),
 ("p","Master of Public Administration, George Mason University, 2015"),
]
JD["federal"] = """Supervisory Management and Program Analyst, GS-0343-13
Specialized Experience: One year of specialized experience equivalent to GS-12, including program evaluation, budget formulation and execution, preparing briefings for senior executives, and supervising staff.
KSAs: knowledge of OMB Circular A-11; skill in data analysis; ability to communicate in writing.
"""

for k, spec in P.items():
    write_docx(os.path.join(R, f"{k}.docx"), spec)
for k, t in JD.items():
    open(os.path.join(J, f"{k}.txt"), "w").write(t)

# PDF versions for a few
for k in ("nurse", "electrician", "sales", "academic"):
    write_pdf(os.path.join(R, f"{k}.pdf"), P[k])

# Designer: two-column table layout with header/footer + portfolio link
d = docx.Document()
d.sections[0].header.paragraphs[0].text = "Lena Park · Product Designer · lena@example.com · lenapark.design · Dribbble: dribbble.com/lenapark"
d.sections[0].footer.paragraphs[0].text = "References available on request"
t = d.add_table(rows=1, cols=2)
left, right = t.rows[0].cells
left.paragraphs[0].add_run("SKILLS").bold = True
for s in ("Figma", "Sketch", "Prototyping", "User research", "Design systems"):
    left.add_paragraph(s)
left.add_paragraph().add_run("EDUCATION").bold = True
left.add_paragraph("BFA Graphic Design, RISD, 2017")
right.paragraphs[0].add_run("EXPERIENCE").bold = True
right.add_paragraph().add_run("Product Designer, Hopper, 2020 – Present").bold = True
right.add_paragraph("Redesigned booking flow for 10M users", style="List Bullet")
right.add_paragraph("Built the Hopper design system in Figma", style="List Bullet")
right.add_paragraph().add_run("Visual Designer, Wieden+Kennedy, 2017 – 2020").bold = True
right.add_paragraph("Campaign visuals for Nike and Airbnb", style="List Bullet")
d.save(os.path.join(R, "designer.docx"))

# Text-box resume: content in a text box (common Canva/Word template)
from docx.oxml import parse_xml
d = docx.Document()
d.add_paragraph().add_run("Sam Taylor").bold = True
d.add_paragraph("sam.taylor@example.com | 555-555-0100")
p = d.add_paragraph()
txbx = parse_xml('''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline><wp:extent cx="5000000" cy="2000000"/><wp:docPr id="1" name="TB"/><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>EXPERIENCE</w:t></w:r></w:p><w:p><w:r><w:t>Office Manager, Bright Dental, 2018 - Present</w:t></w:r></w:p><w:p><w:r><w:t>- Managed scheduling for 4 dentists and billing with Dentrix</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice></mc:AlternateContent></w:r>''')
p._p.append(txbx)
d.add_paragraph().add_run("EDUCATION").bold = True
d.add_paragraph("BA English, State University, 2016")
d.save(os.path.join(R, "textbox.docx"))
print("ok", sorted(os.listdir(R)))
