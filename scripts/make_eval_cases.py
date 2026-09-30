"""Generate the anonymized evaluation cases (P4.2).

Every person, company and school here is fictional. Each case is written
from a ground-truth description, and expected.json is derived from that
same description, so the eval checks what is really in the file rather
than what the parser happens to read.

    python scripts/make_eval_cases.py          # (re)writes data/eval/cases/*

PDF cases are converted with LibreOffice; without it they are skipped.
"""
import json
import os
import sys

import docx
from docx.shared import Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)
OUT = os.path.join(ROOT, "data", "eval", "cases")

LOREM_BOILERPLATE = (
    "About Us\n"
    "Tailspin Health is an equal opportunity employer. We celebrate diversity and are committed to creating an "
    "inclusive environment for all employees. All qualified applicants will receive consideration for employment "
    "without regard to race, color, religion, sex, sexual orientation, gender identity, national origin, disability "
    "or veteran status.\n"
    "Benefits\n"
    "Competitive salary, health insurance, 401(k) matching, flexible hours, remote-friendly culture, learning budget, "
    "wellness stipend, paid parental leave and 25 days of paid time off.\n"
    "Our Values\n"
    "We put patients first. We act with integrity. We win together.\n"
)


# ---------------------------------------------------------------------------
# Ground truth for each case
# ---------------------------------------------------------------------------

CASES = [
    {
        "name": "new-grad",
        "layout": "plain",
        "person": {"name": "Alex Rivera", "email": "alex.rivera@example.com", "phone": "+1 555 010 2233",
                   "location": "Austin, TX", "links": ["linkedin.com/in/alex-rivera-demo"]},
        "summary": "Computer science graduate who builds web APIs in Python and loves clean, tested code.",
        "education": [{"degree": "B.S. in Computer Science", "institution": "Lakeside State University",
                       "dates": "2021 - 2025"}],
        "jobs": [
            {"company": "Contoso Labs", "roles": [("Software Engineering Intern", "Jun 2024", "Aug 2024")],
             "bullets": ["Built a Flask REST API for internal inventory lookups used by 3 teams.",
                         "Wrote pytest suites that raised coverage from 40% to 85%.",
                         "Containerised the service with Docker for the staging environment."]},
        ],
        "projects": [
            {"name": "Campus Ride Share", "bullets": ["Built a React and FastAPI app that matched 200+ riders.",
                                                      "Stored trips in PostgreSQL with indexed queries."]},
        ],
        "skills": {"Languages": ["Python", "JavaScript", "SQL"], "Tools": ["Docker", "Git", "PostgreSQL"]},
        "jd": ("Junior Backend Engineer\nWingtip Software is looking for a Junior Backend Engineer to join our "
               "platform team.\nRequirements\n- Python and a web framework such as FastAPI or Flask\n"
               "- Experience with SQL databases like PostgreSQL\n- Familiarity with Docker and Git\n"
               "- Write automated tests with pytest\nNice to have\n- Kubernetes\n- AWS\n"),
        "expect": {"target_pages": 1, "education_first": True, "jd_company": "Wingtip Software",
                   "must_find": ["Python", "PostgreSQL", "Docker", "Git", "pytest"]},
    },
    {
        "name": "senior-12y",
        "layout": "company_line",
        "person": {"name": "Morgan Blake", "email": "morgan.blake@example.com", "phone": "+44 20 7946 0000",
                   "location": "London, UK", "links": ["github.com/morgan-blake-demo"]},
        "summary": "Engineering leader with 12+ years building payment and data platforms at scale.",
        "education": [{"degree": "M.Sc. in Software Engineering", "institution": "Northgate University",
                       "dates": "2008 - 2010"}],
        "jobs": [
            {"company": "Fabrikam Payments", "location": "London, UK",
             "roles": [("Principal Engineer", "Mar 2020", "Present")],
             "bullets": ["Led a team of 9 engineers rebuilding the card authorisation platform in Java and Kafka.",
                         "Cut p99 authorisation latency from 450 ms to 120 ms.",
                         "Set up SLOs and on-call practices that reduced incidents by 35%.",
                         "Mentored 4 engineers into senior roles."]},
            {"company": "Northwind Traders", "location": "London, UK",
             "roles": [("Senior Software Engineer", "Jan 2016", "Feb 2020")],
             "bullets": ["Designed an event-driven order pipeline on AWS processing 5M orders a day.",
                         "Migrated 30 services from a monolith to Kubernetes.",
                         "Introduced contract testing across 6 teams."]},
            {"company": "Globex Retail", "location": "Manchester, UK",
             "roles": [("Software Engineer", "Aug 2012", "Dec 2015")],
             "bullets": ["Built the Java checkout service handling 2,000 requests per second.",
                         "Automated releases with Jenkins pipelines."]},
            {"company": "Initech", "location": "Leeds, UK",
             "roles": [("Junior Developer", "Sep 2010", "Jul 2012")],
             "bullets": ["Maintained internal reporting tools in Python and SQL.",
                         "Fixed 150+ customer-reported defects."]},
        ],
        "skills": {"Languages": ["Java", "Python", "SQL", "Go"],
                   "Platforms": ["AWS", "Kubernetes", "Kafka", "PostgreSQL"],
                   "Practices": ["SLOs", "Contract testing", "CI/CD"]},
        "jd": ("Staff Engineer, Payments\nUmbrella Fintech is looking for a Staff Engineer to join our payments "
               "group.\nWhat you'll do\n- Lead the design of high-throughput payment services in Java\n"
               "- Own reliability: SLOs, on-call and incident reviews\n- Mentor senior engineers\n"
               "Requirements\n- 10+ years of software engineering experience\n- Deep experience with Kafka "
               "and event-driven systems\n- Production experience with Kubernetes on AWS\n"
               "Nice to have\n- Go\n- PCI DSS\n"),
        "expect": {"target_pages": 2, "education_first": False, "jd_company": "Umbrella Fintech",
                   "must_find": ["Java", "Kafka", "Kubernetes", "AWS"]},
    },
    {
        "name": "career-changer",
        "layout": "plain",
        "person": {"name": "Jamie Chen", "email": "jamie.chen@example.com", "phone": "+61 2 5550 1234",
                   "location": "Sydney, Australia", "links": []},
        "summary": "Former secondary-school maths teacher moving into data analysis; experienced with Excel, "
                   "SQL and presenting numbers to non-experts.",
        "education": [{"degree": "Google Data Analytics Certificate", "institution": "Online",
                       "dates": "2024 - 2025"},
                      {"degree": "B.Ed. in Mathematics Education", "institution": "Harbour University",
                       "dates": "2013 - 2016"}],
        "jobs": [
            {"company": "Bayside High School", "location": "Sydney, Australia",
             "roles": [("Mathematics Teacher", "Feb 2017", "Dec 2024")],
             "bullets": ["Analysed exam results for 600 students in Excel to target revision classes.",
                         "Built a SQL-backed tracker of student progress used by 12 teachers.",
                         "Presented attainment trends to the leadership team each term."]},
        ],
        "projects": [
            {"name": "City Bike Usage Analysis",
             "bullets": ["Cleaned 1M trip records in Python pandas and visualised demand in Tableau."]},
        ],
        "skills": {"Tools": ["Excel", "SQL", "Tableau", "Python", "pandas"]},
        "jd": ("Data Analyst\nContoso Retail is looking for a Data Analyst to join our insights team.\n"
               "Requirements\n- Strong SQL and Excel skills\n- Experience with a BI tool such as Tableau or "
               "Power BI\n- Communicate findings to non-technical stakeholders\n- Python for data cleaning "
               "is a plus\n"),
        "expect": {"target_pages": 1, "education_first": False, "jd_company": "Contoso Retail",
                   "must_find": ["SQL", "Excel", "Tableau", "Python"]},
    },
    {
        "name": "table-docx",
        "layout": "tables",
        "person": {"name": "Riley Novak", "email": "riley.novak@example.com", "phone": "+1 555 010 7788",
                   "location": "Denver, CO", "links": ["linkedin.com/in/riley-novak-demo"]},
        "summary": "Frontend engineer focused on accessible, fast React applications.",
        "education": [{"degree": "B.A. in Interaction Design", "institution": "Mountain View College",
                       "dates": "2014 - 2018"}],
        "jobs": [
            {"company": "Wingtip Toys", "location": "Denver, CO",
             "roles": [("Frontend Engineer", "May 2021", "Present")],
             "bullets": ["Rebuilt the storefront in React and TypeScript, improving Lighthouse scores to 95.",
                         "Led an accessibility audit and fixed 120 WCAG issues.",
                         "Introduced Storybook and visual regression tests."]},
            {"company": "Tailspin Media", "location": "Boulder, CO",
             "roles": [("Web Developer", "Jul 2018", "Apr 2021")],
             "bullets": ["Built marketing pages in Vue.js for 20+ campaigns.",
                         "Cut bundle size by 30% with code splitting."]},
        ],
        "skills": {"Languages": ["TypeScript", "JavaScript", "HTML", "CSS"],
                   "Frameworks": ["React", "Vue.js", "Storybook"]},
        "jd": ("Senior Frontend Engineer\nFabrikam Studios is looking for a Senior Frontend Engineer to join our "
               "web team.\nRequirements\n- 5+ years building web apps with React and TypeScript\n"
               "- Strong grasp of web accessibility (WCAG)\n- Performance tuning and code splitting\n"
               "Nice to have\n- Next.js\n- GraphQL\n"),
        "expect": {"target_pages": 2, "education_first": False, "jd_company": "Fabrikam Studios",
                   "must_find": ["React", "TypeScript", "WCAG"]},
    },
    {
        "name": "promotion-projects",
        "layout": "company_line",
        "person": {"name": "Taylor Brooks", "email": "taylor.brooks@example.com", "phone": "+91 90000 11111",
                   "location": "Pune, India", "links": []},
        "summary": "Data scientist building forecasting and churn models on Spark and Databricks.",
        "education": [{"degree": "B.Tech in Electrical Engineering", "institution": "Riverside Institute of "
                       "Technology", "dates": "2014 - 2018"}],
        "jobs": [
            {"company": "Northwind Analytics", "location": "Pune, India",
             "roles": [("Senior Data Scientist", "Apr 2022", "Present"), ("Data Scientist", "Jul 2018", "Mar 2022")],
             "groups": [
                 ("Demand Forecasting Platform",
                  ["Built LightGBM demand forecasts on Databricks for 4,000 stores.",
                   "Cut forecast error (MAPE) from 18% to 11%."]),
                 ("Customer Churn Model",
                  ["Trained an XGBoost churn model on PySpark features for 2M customers.",
                   "Deployed it with MLflow and weekly retraining."]),
             ]},
        ],
        "skills": {"Languages": ["Python", "SQL", "PySpark"], "ML": ["LightGBM", "XGBoost", "MLflow"],
                   "Platforms": ["Databricks", "Azure"]},
        "jd": ("Machine Learning Engineer\nGlobex AI is looking for a Machine Learning Engineer to join our "
               "forecasting team.\nRequirements\n- Python and SQL\n- Gradient boosting models such as "
               "LightGBM or XGBoost\n- Spark or Databricks for large-scale training\n- MLflow or similar "
               "for model tracking\nNice to have\n- Kubernetes\n"),
        "expect": {"target_pages": 2, "education_first": False, "jd_company": "Globex AI",
                   "must_find": ["Python", "LightGBM", "XGBoost", "Databricks", "MLflow"]},
    },
    {
        "name": "pdf-mid",
        "layout": "company_line",
        "format": "pdf",
        "person": {"name": "Casey Morgan", "email": "casey.morgan@example.com", "phone": "+49 30 5550 9876",
                   "location": "Berlin, Germany", "links": ["github.com/casey-morgan-demo"]},
        "summary": "DevOps engineer automating cloud infrastructure with Terraform and Kubernetes.",
        "education": [{"degree": "B.Sc. in Computer Science", "institution": "Spree Technical University",
                       "dates": "2012 - 2016"}],
        "jobs": [
            {"company": "Contoso Cloud", "location": "Berlin, Germany",
             "roles": [("DevOps Engineer", "Jan 2020", "Present")],
             "bullets": ["Managed 40 AWS accounts with Terraform modules.",
                         "Ran Kubernetes clusters serving 300 microservices.",
                         "Built GitHub Actions pipelines that cut deploy time from 40 to 8 minutes."]},
            {"company": "Fabrikam Hosting", "location": "Hamburg, Germany",
             "roles": [("Systems Administrator", "Sep 2016", "Dec 2019")],
             "bullets": ["Automated Linux server provisioning with Ansible.",
                         "Set up Prometheus and Grafana monitoring for 500 hosts."]},
        ],
        "skills": {"Cloud": ["AWS", "Terraform", "Kubernetes"], "Tooling": ["Ansible", "Prometheus", "Grafana"]},
        "jd": ("Site Reliability Engineer\nNorthwind Cloud is looking for a Site Reliability Engineer to join our "
               "platform team.\nRequirements\n- Infrastructure as code with Terraform\n- Kubernetes in "
               "production\n- Monitoring with Prometheus and Grafana\n- CI/CD pipelines\nNice to have\n"
               "- Go\n"),
        "expect": {"target_pages": 2, "education_first": False, "jd_company": "Northwind Cloud",
                   "must_find": ["Terraform", "Kubernetes", "Prometheus", "Grafana"]},
    },
    {
        "name": "boilerplate-jd",
        "layout": "plain",
        "person": {"name": "Jordan Ellis", "email": "jordan.ellis@example.com", "phone": "+1 555 010 4455",
                   "location": "Chicago, IL", "links": []},
        "summary": "Healthcare data engineer building HIPAA-compliant pipelines.",
        "education": [{"degree": "B.S. in Information Systems", "institution": "Lakeshore University",
                       "dates": "2013 - 2017"}],
        "jobs": [
            {"company": "Tailspin Clinics", "location": "Chicago, IL",
             "roles": [("Data Engineer", "Jun 2019", "Present")],
             "bullets": ["Built Airflow pipelines loading 2TB of claims data into Snowflake daily.",
                         "Implemented HIPAA-compliant de-identification for patient records.",
                         "Modelled the analytics layer in dbt for 15 analysts."]},
            {"company": "Initech Health", "location": "Chicago, IL",
             "roles": [("ETL Developer", "Jul 2017", "May 2019")],
             "bullets": ["Wrote SQL Server ETL jobs for billing reports.",
                         "Cut nightly load time by 50%."]},
        ],
        "skills": {"Data": ["SQL", "Python", "Airflow", "dbt", "Snowflake"]},
        "jd": ("Senior Data Engineer\nTailspin Health is looking for a Senior Data Engineer to join our data "
               "platform team.\n" + LOREM_BOILERPLATE + "Requirements\n- Build batch pipelines with Airflow\n"
               "- Snowflake and dbt experience\n- Healthcare data and HIPAA knowledge\n- Strong SQL and "
               "Python\nNice to have\n- Kafka\n" + LOREM_BOILERPLATE),
        "expect": {"target_pages": 2, "education_first": False, "jd_company": "Tailspin Health",
                   "must_find": ["Airflow", "Snowflake", "dbt", "HIPAA", "SQL", "Python"],
                   "must_not_be_keywords": ["401(k)", "diversity", "wellness"]},
    },
    {
        "name": "product-manager",
        "layout": "plain",
        "person": {"name": "Avery Quinn", "email": "avery.quinn@example.com", "phone": "+1 555 010 6677",
                   "location": "Seattle, WA", "links": ["linkedin.com/in/avery-quinn-demo"]},
        "summary": "Product manager shipping B2B SaaS features from discovery to launch.",
        "education": [{"degree": "MBA", "institution": "Cascade Business School", "dates": "2015 - 2017"}],
        "jobs": [
            {"company": "Fabrikam Software", "location": "Seattle, WA",
             "roles": [("Senior Product Manager", "Mar 2021", "Present")],
             "bullets": ["Owned the roadmap for a billing product with $12M ARR.",
                         "Ran 40 customer discovery interviews that reshaped the onboarding flow.",
                         "Launched usage-based pricing with A/B tests that lifted conversion by 9%."]},
            {"company": "Contoso SaaS", "location": "Bellevue, WA",
             "roles": [("Product Manager", "Aug 2017", "Feb 2021")],
             "bullets": ["Wrote PRDs and user stories for a team of 8 engineers.",
                         "Defined OKRs and a metrics dashboard in Looker."]},
        ],
        "skills": {"Product": ["Roadmapping", "Discovery", "A/B testing", "OKRs"], "Tools": ["Jira", "Looker", "SQL"]},
        "jd": ("Group Product Manager\nWingtip Cloud is looking for a Group Product Manager to join our "
               "platform organisation.\nRequirements\n- 6+ years of product management in B2B SaaS\n"
               "- Customer discovery and roadmapping\n- Data-driven decisions with A/B testing and SQL\n"
               "- Excellent communication and stakeholder management\nNice to have\n- Pricing and "
               "monetisation experience\n"),
        "expect": {"target_pages": 2, "education_first": False, "jd_company": "Wingtip Cloud",
                   "must_find": ["SQL", "A/B testing"]},
    },
]


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

def _heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)


def _bullets(doc, items):
    for text in items:
        doc.add_paragraph(text, style="List Bullet")


def _job_lines(doc, job, layout):
    roles = job["roles"]
    if layout == "company_line":
        p = doc.add_paragraph()
        p.add_run(job["company"]).bold = True
        if job.get("location"):
            p.add_run("\t" + job["location"])
        for title, start, end in roles:
            t = doc.add_paragraph()
            t.add_run(title).italic = True
            t.add_run(f"\t{start} - {end}")
    else:  # "plain": "Title | Company | dates"
        for title, start, end in roles:
            p = doc.add_paragraph()
            p.add_run(f"{title} | {job['company']} | {start} - {end}").bold = True
    for group, items in job.get("groups", []):
        g = doc.add_paragraph()
        g.add_run(group).bold = True
        _bullets(doc, items)
    _bullets(doc, job.get("bullets", []))


def write_docx(case, path):
    doc = docx.Document()
    person = case["person"]
    name = doc.add_paragraph()
    name_run = name.add_run(person["name"])
    name_run.bold = True
    name_run.font.size = Pt(20)
    contact = " | ".join([person["email"], person["phone"], person["location"], *person["links"]])
    if case["layout"] == "tables":
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = f"{person['email']} | {person['phone']}"
        table.cell(0, 1).text = " | ".join([person["location"], *person["links"]])
    else:
        doc.add_paragraph(contact)

    _heading(doc, "Summary")
    doc.add_paragraph(case["summary"])
    _heading(doc, "Experience")
    for job in case["jobs"]:
        _job_lines(doc, job, "company_line" if case["layout"] == "tables" else case["layout"])
    for project in case.get("projects", []):
        pass
    if case.get("projects"):
        _heading(doc, "Projects")
        for project in case["projects"]:
            p = doc.add_paragraph()
            p.add_run(project["name"]).bold = True
            _bullets(doc, project["bullets"])
    _heading(doc, "Skills")
    if case["layout"] == "tables":
        table = doc.add_table(rows=len(case["skills"]), cols=2)
        for i, (category, items) in enumerate(case["skills"].items()):
            table.cell(i, 0).text = category
            table.cell(i, 1).text = ", ".join(items)
    else:
        for category, items in case["skills"].items():
            p = doc.add_paragraph()
            p.add_run(f"{category}: ").bold = True
            p.add_run(", ".join(items))
    _heading(doc, "Education")
    for edu in case["education"]:
        p = doc.add_paragraph()
        p.add_run(edu["institution"]).bold = True
        e = doc.add_paragraph(f"{edu['degree']}\t{edu['dates']}")
    doc.save(path)


def expected_for(case):
    jobs = case["jobs"]
    bullets = sum(len(j.get("bullets", [])) + sum(len(g[1]) for g in j.get("groups", [])) for j in jobs)
    bullets += sum(len(p["bullets"]) for p in case.get("projects", []))
    return {
        "name": case["person"]["name"],
        "email": case["person"]["email"],
        "jobs": len(jobs),
        "companies": [j["company"] for j in jobs],
        "roles": [[r[0] for r in j["roles"]] for j in jobs],
        "bullets": bullets,
        "projects": len(case.get("projects", [])),
        "skills": sorted({s for items in case["skills"].values() for s in items}),
        "education": len(case["education"]),
        **case["expect"],
    }


def main():
    from app.rendering.pdf_converter import PdfConverter

    manifest = []
    for case in CASES:
        folder = os.path.join(OUT, case["name"])
        os.makedirs(folder, exist_ok=True)
        docx_path = os.path.join(folder, "resume.docx")
        write_docx(case, docx_path)
        resume = docx_path
        if case.get("format") == "pdf":
            pdf = PdfConverter().convert_docx_to_pdf(docx_path, folder)
            if not pdf:
                print(f"skip {case['name']}: LibreOffice not available for the PDF")
                continue
            os.remove(docx_path)
            resume = pdf
        with open(os.path.join(folder, "jd.txt"), "w", encoding="utf-8") as f:
            f.write(case["jd"])
        with open(os.path.join(folder, "expected.json"), "w", encoding="utf-8") as f:
            json.dump(expected_for(case), f, indent=2, ensure_ascii=False)
        rel = lambda p: os.path.relpath(p, ROOT)
        manifest.append({"name": case["name"], "resume": rel(resume), "jd": rel(os.path.join(folder, "jd.txt")),
                         "expected": rel(os.path.join(folder, "expected.json"))})
        print(f"wrote {case['name']}")
    with open(os.path.join(OUT, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)


if __name__ == "__main__":
    main()
